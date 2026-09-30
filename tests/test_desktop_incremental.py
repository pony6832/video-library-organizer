import os
import uuid

import pytest
from openpyxl import load_workbook

from media_catalog.bootstrap import bootstrap_workspace
from media_catalog.batch_analysis import analyze_pending
from media_catalog.database import CatalogDatabase
from media_catalog.inference import Analysis


def test_catalog_writes_visible_excel_at_first_100_then_throttles(tmp_path, monkeypatch):
    from media_catalog import bootstrap

    for number in range(205):
        (tmp_path / f'clip-{number:03}.mp4').write_bytes(b'video')
    rows_at_write = []
    real_write = bootstrap.write_excel

    def capture(records, path):
        rows_at_write.append(len(list(records)))
        return real_write(records, path)

    monkeypatch.setattr(bootstrap, 'write_excel', capture)
    result = bootstrap_workspace(tmp_path, video_only=True)
    # Later 100-item checkpoints are throttled to one per minute, so a fast
    # scan does not rewrite the whole workbook quadratically.
    assert rows_at_write == [100, 205]
    book = load_workbook(result.workspace.excel_path, read_only=True)
    try:
        assert sum(1 for _ in book.active.iter_rows(values_only=True)) == 206
    finally:
        book.close()


def test_analysis_syncs_every_100_and_on_final_remainder(tmp_path):
    from media_catalog import batch_analysis

    for number in range(105):
        (tmp_path / f'clip-{number:03}.mp4').write_bytes(b'video')
    workspace = bootstrap_workspace(tmp_path, video_only=True).workspace
    writes = []

    class Local:
        def analyze(self, path):
            return Analysis('影片內容', ('重點',), ('關鍵字',))

    def write(records, path):
        writes.append(sum(bool(record.description) for record in records))
        return batch_analysis.write_excel(records, path)

    result = analyze_pending(workspace, Local(), excel_writer=write, video_only=True)
    assert result.analyzed == 105
    assert writes == [100, 105]
    assert sum(bool(r.description) for r in CatalogDatabase(workspace.database_path).list_records()) == 105


def test_progress_text_includes_percent_and_count():
    from media_catalog.desktop import analysis_progress_text

    assert analysis_progress_text(23, '23 / 100', '77') == '23%　已完成 23 / 100　未完成 77'
    assert analysis_progress_text(0, '清冊已收錄 5 部影片；尚未開始本次分析', '啟動分析後計算').startswith('清冊已收錄')


def test_cataloging_view_reports_written_checkpoint(tmp_path):
    from unittest.mock import Mock
    from media_catalog.desktop import DesktopApplication
    from media_catalog.supervisor import SupervisorSnapshot

    (tmp_path / 'clip.mp4').write_bytes(b'video')
    workspace = bootstrap_workspace(tmp_path, video_only=True).workspace
    app = DesktopApplication.__new__(DesktopApplication)
    app.workspace = workspace
    app.media_root = tmp_path
    app.cloud_var = Mock()
    model = app._view_model(SupervisorSnapshot('cataloging', True, None))
    assert model.video_count == 1
    assert '已寫入 1' in model.progress_text


@pytest.mark.skipif(os.name != 'nt', reason='Windows Credential Manager')
def test_credential_roundtrip_replaces_without_plaintext_file():
    from media_catalog.credential_store import read_key, write_key, delete_key

    target = f'MediaCatalogVideoDesktop/test-{uuid.uuid4()}'
    try:
        assert read_key(target) is None
        write_key('test-first-key', target)
        assert read_key(target) == 'test-first-key'
        write_key('test-replaced-key', target)
        assert read_key(target) == 'test-replaced-key'
    finally:
        delete_key(target)
    assert read_key(target) is None
