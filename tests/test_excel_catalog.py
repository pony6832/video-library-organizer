import os
from itertools import zip_longest
from pathlib import Path


def test_strict_review_reader_normalizes_malformed_xlsx_zip(tmp_path):
    import pytest
    from zipfile import ZipFile
    from media_catalog.excel_catalog import ReviewedPathsError, read_reviewed_paths_strict

    source = tmp_path / "malformed.xlsx"
    with ZipFile(source, "w") as archive:
        archive.writestr("unrelated.txt", "not an Excel package")
    original = source.read_bytes()
    with pytest.raises(ReviewedPathsError):
        read_reviewed_paths_strict(source)
    assert source.read_bytes() == original

import pytest
from openpyxl import load_workbook

from media_catalog.database import CatalogDatabase
from media_catalog.excel_catalog import (
    CATALOG_HEADERS,
    ReviewedPathsError,
    read_reviewed_paths,
    write_excel,
)


def test_write_excel_uses_spec_columns_and_persisted_values(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    photo_path = tmp_path / "旅行照片.jpg"
    photo_path.write_bytes(b"photo")
    record = database.upsert_discovered(photo_path, "abc", "image/jpeg")
    output = tmp_path / "媒體清冊.xlsx"

    saved_path = write_excel([record], output)

    workbook = load_workbook(saved_path, read_only=True)
    sheet = workbook["媒體清冊"]
    assert [cell.value for cell in sheet[1]] == list(CATALOG_HEADERS)
    assert sheet.cell(2, 1).value == "待處理"
    assert sheet.cell(2, 2).value == "旅行照片.jpg"
    assert sheet.cell(2, 3).value == str(photo_path.resolve())
    assert sheet.cell(2, 4).value == "照片 (JPEG)"
    workbook.close()


def test_write_excel_labels_analyzed_records_for_review(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    photo_path = tmp_path / "賀卡.png"
    photo_path.write_bytes(b"photo")
    record = database.upsert_discovered(photo_path, "abc", "image/png")
    analyzed = database.save_analysis(
        record.id,
        description="紅色馬年賀卡。",
        highlights=("金色馬",),
        keywords=("賀卡", "馬年"),
    )

    saved_path = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path, read_only=True)
    sheet = workbook["媒體清冊"]
    assert sheet.cell(2, 1).value == "待確認"
    assert sheet.cell(2, 5).value == "紅色馬年賀卡。"
    assert sheet.cell(2, 9).value is not None
    workbook.close()


def test_write_excel_adds_review_dropdown_to_status_rows(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    first_path = tmp_path / "first.png"
    second_path = tmp_path / "second.png"
    first_path.write_bytes(b"first")
    second_path.write_bytes(b"second")
    first = database.upsert_discovered(first_path, "first", "image/png")
    second = database.upsert_discovered(second_path, "second", "image/png")

    saved_path = write_excel([first, second], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path)
    sheet = workbook["媒體清冊"]
    validations = list(sheet.data_validations.dataValidation)
    assert len(validations) == 1
    validation = validations[0]
    assert validation.type == "list"
    assert validation.formula1 == '"待確認,已審核"'
    assert str(validation.sqref) == "A2:A3"
    assert validation.showDropDown is False
    workbook.close()


def test_write_excel_omits_review_dropdown_for_empty_catalog(
    tmp_path: Path,
) -> None:
    saved_path = write_excel([], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(saved_path)
    sheet = workbook["媒體清冊"]
    assert list(sheet.data_validations.dataValidation) == []
    assert sheet.max_row == 1
    workbook.close()


def test_write_excel_preserves_existing_workbook_when_atomic_replace_fails(
    tmp_path: Path, monkeypatch
) -> None:
    output = tmp_path / "媒體清冊.xlsx"
    write_excel([], output)
    original = output.read_bytes()

    def locked_replace(*_args, **_kwargs):
        raise PermissionError("workbook is open")

    monkeypatch.setattr(os, "replace", locked_replace)

    with pytest.raises(PermissionError, match="open"):
        write_excel([], output)

    assert output.read_bytes() == original
    assert list(tmp_path.glob(".媒體清冊.*.tmp.xlsx")) == []


def test_path_cell_links_to_existing_source(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "片段 01.mp4"
    source.write_bytes(b"video")
    record = database.upsert_discovered(source, "video-1", "video/mp4")

    output = write_excel([record], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output)
    path_cell = workbook["媒體清冊"].cell(2, 3)
    assert path_cell.value == str(source.resolve())
    assert path_cell.hyperlink is not None
    assert path_cell.hyperlink.target == source.resolve().as_uri()
    assert path_cell.style == "Hyperlink"
    workbook.close()


def test_path_cell_stays_plain_when_source_is_missing(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "missing.mp4"
    record = database.upsert_discovered(source, "missing", "video/mp4")

    output = write_excel([record], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output)
    path_cell = workbook["媒體清冊"].cell(2, 3)
    assert path_cell.value == str(source.resolve())
    assert path_cell.hyperlink is None
    assert path_cell.style != "Hyperlink"
    workbook.close()


def test_atomic_rebuild_preserves_reviewed_status_by_path(
    tmp_path: Path,
) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "reviewed.png"
    source.write_bytes(b"photo")
    record = database.upsert_discovered(source, "reviewed", "image/png")
    analyzed = database.save_analysis(
        record.id,
        description="已完成描述",
        highlights=("清晰",),
        keywords=("照片",),
    )
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    workbook = load_workbook(output)
    workbook["媒體清冊"].cell(2, 1).value = "已審核"
    workbook.save(output)
    workbook.close()

    write_excel([analyzed], output)

    workbook = load_workbook(output)
    assert workbook["媒體清冊"].cell(2, 1).value == "已審核"
    workbook.close()
    assert read_reviewed_paths(output) == {str(source.resolve())}


def test_gemini_warning_keeps_analysis_visible_for_manual_review(
    tmp_path: Path,
) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "photo.jpg"
    source.write_bytes(b"photo")
    record = database.upsert_discovered(source, "photo", "image/jpeg")
    analyzed = database.save_analysis(
        record.id,
        description="本地分析仍可使用。",
        highlights=("保留重點",),
        keywords=("本地備援",),
        warning="Gemini 強化失敗:GeminiError",
    )

    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")

    workbook = load_workbook(output, read_only=True)
    row = tuple(
        cell.value for cell in next(workbook["媒體清冊"].iter_rows(min_row=2))
    )
    workbook.close()
    assert row[0] == "待確認"
    assert row[4:7] == (
        "本地分析仍可使用。",
        "保留重點",
        "本地備援",
    )
    assert row[11] == "Gemini 強化失敗:GeminiError"


def test_write_excel_accepts_stream_without_materializing_records(tmp_path: Path, monkeypatch) -> None:
    from media_catalog import excel_catalog
    from media_catalog.database import CatalogDatabase
    database = CatalogDatabase(tmp_path / 'catalog.sqlite')
    records = []
    for index in range(3):
        source = tmp_path / f'clip-{index}.mp4'
        source.write_bytes(b'clip')
        records.append(database.upsert_discovered(source, str(index), 'video/mp4'))

    # This checks the implementation uses write-only mode; output is checked too.
    original = excel_catalog.Workbook
    modes = []
    def recording_workbook(*args, **kwargs):
        modes.append(kwargs.get('write_only'))
        return original(*args, **kwargs)
    monkeypatch.setattr(excel_catalog, 'Workbook', recording_workbook)
    output = write_excel((record for record in records), tmp_path / 'catalog.xlsx')
    assert modes == [True]
    workbook = load_workbook(output)
    assert workbook.active.max_row == 4
    assert str(workbook.active.data_validations.dataValidation[0].sqref) == 'A2:A4'
    workbook.close()


def test_corrupt_existing_catalog_is_not_overwritten(tmp_path: Path) -> None:
    output = tmp_path / "catalog.xlsx"
    output.write_bytes(b"damaged-review-data")
    with pytest.raises(ReviewedPathsError):
        write_excel([], output)
    assert output.read_bytes() == b"damaged-review-data"


def test_model_text_is_literal_not_excel_formula(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    record = database.upsert_discovered(tmp_path / "photo.jpg", "a", "image/jpeg")
    record = database.save_analysis(record.id, description='=HYPERLINK("https://example.com","click")', highlights=("=1+1",), keywords=("=2+2",))
    output = write_excel([record], tmp_path / "catalog.xlsx")
    workbook = load_workbook(output)
    try:
        assert workbook.active.cell(2, 5).data_type == "s"
        assert workbook.active.cell(2, 6).data_type == "s"
        assert workbook.active.cell(2, 7).data_type == "s"
    finally:
        workbook.close()


def test_review_status_survives_windows_path_case_changes(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    record = database.upsert_discovered(tmp_path / "photo.jpg", "a", "image/jpeg")
    output = write_excel([record], tmp_path / "catalog.xlsx")
    workbook = load_workbook(output)
    workbook.active.cell(2, 1).value = "已審核"
    workbook.active.cell(2, 3).value = str(record.path).upper()
    workbook.save(output)
    workbook.close()
    write_excel([record], output)
    workbook = load_workbook(output)
    try:
        assert workbook.active.cell(2, 1).value == "已審核"
    finally:
        workbook.close()


def _analyzed_record(tmp_path: Path, name: str = "clip.mp4", description: str = "模型描述"):
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / name
    source.write_bytes(b"clip")
    record = database.upsert_discovered(source, name, "video/mp4")
    analyzed = database.save_analysis(
        record.id, description=description, highlights=("重點一",), keywords=("關鍵字",)
    )
    return database, analyzed


def _edit_cells(output: Path, **cells) -> None:
    workbook = load_workbook(output)
    sheet = workbook["媒體清冊"]
    headers = [cell.value for cell in sheet[1]]
    for header, value in cells.items():
        if header not in headers:
            headers.append(header)
            sheet.cell(1, len(headers)).value = header
        sheet.cell(2, headers.index(header) + 1).value = value
    workbook.save(output)
    workbook.close()


def _row(output: Path) -> dict:
    workbook = load_workbook(output, read_only=True)
    try:
        rows = list(workbook["媒體清冊"].iter_rows(values_only=True))
    finally:
        workbook.close()
    return dict(zip_longest(rows[0], rows[1]))


def test_hand_edits_survive_rebuild_and_baseline_sheet_is_hidden(tmp_path: Path) -> None:
    from media_catalog.excel_catalog import BASELINE_SHEET

    _, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    _edit_cells(output, 內容描述="人工修正的描述", 關鍵字="人工、標籤", 拍攝時間="2026-09-01")

    write_excel([analyzed], output)
    write_excel([analyzed], output)

    row = _row(output)
    assert row["內容描述"] == "人工修正的描述"
    assert row["關鍵字"] == "人工、標籤"
    assert row["拍攝時間"] == "2026-09-01"
    assert row["重點"] == "重點一"
    workbook = load_workbook(output)
    baseline = workbook[BASELINE_SHEET]
    assert baseline.sheet_state == "hidden"
    # The media library panel finds catalog sheets by these headers.
    assert not {"完整路徑", "檔名"} & {cell.value for cell in baseline[1]}
    workbook.close()


def test_new_analysis_updates_untouched_fields_but_not_hand_edits(tmp_path: Path) -> None:
    database, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    _edit_cells(output, 內容描述="人工修正的描述")

    updated = database.save_analysis(
        analyzed.id, description="Gemini 新描述", highlights=("新重點",), keywords=("新關鍵字",)
    )
    write_excel([updated], output)

    row = _row(output)
    assert row["內容描述"] == "人工修正的描述"
    assert row["重點"] == "新重點"
    assert row["關鍵字"] == "新關鍵字"


def test_blank_cells_and_own_status_labels_are_refilled(tmp_path: Path) -> None:
    _, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    _edit_cells(output, 狀態="處理中", 重點=None)

    write_excel([analyzed], output)

    row = _row(output)
    assert (row["狀態"], row["重點"]) == ("待確認", "重點一")


def test_columns_added_by_people_are_carried_over_by_path(tmp_path: Path) -> None:
    _, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    _edit_cells(output, 挑選="是", 評等=4, 備註="人工備註")

    write_excel([analyzed], output)

    row = _row(output)
    assert (row["挑選"], row["評等"], row["備註"]) == ("是", 4, "人工備註")
    workbook = load_workbook(output)
    sheet = workbook["媒體清冊"]
    assert sheet.auto_filter.ref == "A1:O2"
    assert sheet.cell(2, 15).data_type == "s"
    workbook.close()


def test_legacy_workbook_keeps_edits_made_after_the_last_analysis(tmp_path: Path) -> None:
    from media_catalog.excel_catalog import BASELINE_SHEET

    database, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    workbook = load_workbook(output)
    del workbook[BASELINE_SHEET]  # as written by 0.3.0 Beta 1 and earlier
    workbook["媒體清冊"].cell(2, 5).value = "人工修正的描述"
    workbook.save(output)
    workbook.close()

    write_excel([analyzed], output)

    assert _row(output)["內容描述"] == "人工修正的描述"


def test_legacy_workbook_yields_to_analysis_newer_than_the_file(tmp_path: Path) -> None:
    from media_catalog.excel_catalog import BASELINE_SHEET

    database, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    workbook = load_workbook(output)
    del workbook[BASELINE_SHEET]
    workbook.save(output)
    workbook.close()
    old = output.stat().st_mtime - 3600
    os.utime(output, (old, old))

    updated = database.save_analysis(
        analyzed.id, description="Gemini 新描述", highlights=("新重點",), keywords=("新關鍵字",)
    )
    write_excel([updated], output)

    assert _row(output)["內容描述"] == "Gemini 新描述"


def test_edit_saved_during_rebuild_is_not_lost(tmp_path: Path, monkeypatch) -> None:
    from media_catalog import excel_catalog

    _, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    original = excel_catalog._read_existing
    calls = []

    def read_then_edit(source):
        existing = original(source)
        calls.append(source)
        if len(calls) == 1:
            _edit_cells(output, 內容描述="重建途中的人工修改")
            os.utime(output, ns=(existing.stamp[0] + 10**9, existing.stamp[0] + 10**9))
        return existing

    monkeypatch.setattr(excel_catalog, "_read_existing", read_then_edit)
    write_excel([analyzed], output)

    assert len(calls) == 2
    assert _row(output)["內容描述"] == "重建途中的人工修改"


def test_status_set_by_hand_survives_until_changed_again(tmp_path: Path) -> None:
    database, analyzed = _analyzed_record(tmp_path)
    output = write_excel([analyzed], tmp_path / "媒體清冊.xlsx")
    _edit_cells(output, 狀態="需複查")

    from media_catalog.models import Status

    completed = database.set_status(analyzed.id, Status.COMPLETED)
    write_excel([completed], output)
    assert _row(output)["狀態"] == "需複查"

    _edit_cells(output, 狀態="已審核")
    write_excel([completed], output)
    assert _row(output)["狀態"] == "已審核"
    assert read_reviewed_paths(output) == {str(completed.path.resolve())}


def test_status_follows_analysis_when_not_edited(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    source = tmp_path / "clip.mp4"
    source.write_bytes(b"clip")
    pending = database.upsert_discovered(source, "clip", "video/mp4")
    output = write_excel([pending], tmp_path / "媒體清冊.xlsx")
    assert _row(output)["狀態"] == "待處理"

    analyzed = database.save_analysis(
        pending.id, description="描述", highlights=("重點",), keywords=("關鍵字",)
    )
    write_excel([analyzed], output)
    assert _row(output)["狀態"] == "待確認"


def test_row_appended_by_another_program_keeps_its_edits(tmp_path: Path) -> None:
    database, first = _analyzed_record(tmp_path, "a.mp4")
    output = write_excel([first], tmp_path / "媒體清冊.xlsx")
    second_path = tmp_path / "b.mp4"
    second_path.write_bytes(b"clip")
    second = database.save_analysis(
        database.upsert_discovered(second_path, "b", "video/mp4").id,
        description="模型描述 b", highlights=("重點 b",), keywords=("關鍵字 b",),
    )
    # The media library panel appends a row for a video it already shows.
    workbook = load_workbook(output)
    sheet = workbook["媒體清冊"]
    sheet.append(["已確認", "b.mp4", str(second_path.resolve()), "影片 (MP4)", "人工描述 b", "重點 b", "關鍵字 b"])
    workbook.save(output)
    workbook.close()

    write_excel([first, second], output)

    workbook = load_workbook(output, read_only=True)
    rows = {row[2]: row for row in workbook["媒體清冊"].iter_rows(min_row=2, values_only=True)}
    workbook.close()
    assert len(rows) == 2
    row = rows[str(second_path.resolve())]
    assert (row[0], row[4], row[5]) == ("已確認", "人工描述 b", "重點 b")

    # Once we have written the row, later analysis updates the untouched cells.
    updated = database.save_analysis(second.id, description="新描述", highlights=("新重點",), keywords=("新",))
    write_excel([first, updated], output)
    workbook = load_workbook(output, read_only=True)
    rows = {row[2]: row for row in workbook["媒體清冊"].iter_rows(min_row=2, values_only=True)}
    workbook.close()
    row = rows[str(second_path.resolve())]
    assert (row[0], row[4], row[5]) == ("已確認", "人工描述 b", "新重點")
