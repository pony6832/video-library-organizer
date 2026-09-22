from __future__ import annotations

import hashlib
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from media_catalog.batch_analysis import analyze_pending
from media_catalog.bootstrap import bootstrap_workspace
from media_catalog.frame_selector import FrameSelector
from media_catalog.inference import Analysis
from media_catalog.run_state import RunStateStore
from media_catalog.scene_segments import SceneSegmenter
from media_catalog.segment_pipeline import SegmentPipeline


LOCAL_IMAGE = Analysis(
    "測試圖像顯示單色畫面。",
    ("圖像資料已讀取",),
    ("測試圖像",),
)
WEAK_VIDEO = Analysis("一個場景。", ("畫面",), ("內容",))
STRONG_VIDEO = Analysis(
    "測試影片由紅色與藍色兩個場景組成。",
    ("兩個明顯色彩場景",),
    ("測試影片", "場景轉換"),
)


class FakeLocalAnalyzer:
    def analyze(self, _source: Path) -> Analysis:
        return LOCAL_IMAGE

    def analyze_frames(
        self, _frames: tuple[Path, ...], *, ocr_text: str = ""
    ) -> Analysis:
        return WEAK_VIDEO

    def summarize_segments(
        self, _analyses: tuple[Analysis, ...]
    ) -> Analysis:
        return STRONG_VIDEO


class FakeGemini:
    is_configured = True
    discovery_error = None

    def discover_model(self):
        return 'gemini-synthetic'

    def analyze(self, _request) -> Analysis:
        return STRONG_VIDEO


class FakeRuntime:
    def __init__(self, workspace) -> None:
        self.run_state = RunStateStore(
            workspace.database_path, excel_path=workspace.excel_path
        )
        self.local_analyzer = FakeLocalAnalyzer()
        self.segment_pipeline = SegmentPipeline(
            segmenter=SceneSegmenter(timeout_seconds=30),
            selector=FrameSelector(),
            local_analyzer=self.local_analyzer,
            gemini_client=FakeGemini(),
            store=self.run_state,
            output_root=workspace.temp_dir / "a-plus-smoke-segments",
        )

    def analyze(self, source: Path) -> Analysis:
        return self.local_analyzer.analyze(source)


@pytest.mark.skipif(
    shutil.which("ffmpeg") is None or shutil.which("ffprobe") is None,
    reason="FFmpeg integration tools are unavailable",
)
def test_headless_a_plus_pipeline_preserves_sources_and_builds_excel(
    tmp_path: Path,
) -> None:
    root = tmp_path / "media"
    root.mkdir()
    image_path = root / "sample.jpg"
    video_path = root / "scenes.mp4"
    Image.new("RGB", (96, 64), "navy").save(image_path)
    subprocess.run(
        [
            "ffmpeg",
            "-hide_banner",
            "-loglevel",
            "error",
            "-f",
            "lavfi",
            "-i",
            "color=c=red:s=160x120:d=2:r=10",
            "-f",
            "lavfi",
            "-i",
            "color=c=blue:s=160x120:d=2:r=10",
            "-filter_complex",
            "[0:v][1:v]concat=n=2:v=1:a=0,format=yuv420p[v]",
            "-map",
            "[v]",
            "-y",
            str(video_path),
        ],
        capture_output=True,
        check=True,
        timeout=30,
    )
    hashes = {
        path: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (image_path, video_path)
    }
    workspace = bootstrap_workspace(root).workspace

    result = analyze_pending(workspace, FakeRuntime(workspace))

    assert result.failed == 0
    assert result.remaining == 0
    assert result.excel_sync_pending is False
    workbook = load_workbook(workspace.excel_path, read_only=False)
    try:
        sheet = workbook["媒體清冊"]
        assert sheet.max_row == 3
        for row in range(2, 4):
            assert sheet.cell(row, 3).hyperlink is not None
            assert str(sheet.cell(row, 5).value).strip()
            assert str(sheet.cell(row, 6).value).strip()
            assert str(sheet.cell(row, 7).value).strip()
    finally:
        workbook.close()
    assert {
        path: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in hashes
    } == hashes
