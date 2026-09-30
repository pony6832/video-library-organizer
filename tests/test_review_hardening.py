"""Regression tests for robustness and scale fixes (review 2026-10)."""
from __future__ import annotations

import os
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from media_catalog import excel_catalog, scanner as scanner_module
from media_catalog import setup_environment
from media_catalog.bootstrap import bootstrap_workspace
from media_catalog.database import CatalogDatabase
from media_catalog.excel_catalog import ExcelCheckpoint, try_write_excel, write_excel
from media_catalog.models import Status
from media_catalog.process_utils import credential_free_environment
from media_catalog.run_state import RunStateStore
from media_catalog.status_ui import StatusViewModel
from media_catalog.supervisor import WorkerSupervisor

from test_supervisor import FakeProcess, ProcessFactory, prepared_root


# --- supervisor -----------------------------------------------------------

def test_exit_code_3_is_incomplete_not_a_crash(tmp_path: Path) -> None:
    root, _, store, _ = prepared_root(tmp_path)
    factory = ProcessFactory([FakeProcess(returncode=3)])
    supervisor = WorkerSupervisor(
        process_factory=factory, store_factory=lambda _workspace: store
    )

    supervisor.start(root, tmp_path / "skill")
    snapshot = supervisor.poll()

    assert snapshot.status == "incomplete"
    assert snapshot.worker_alive is False
    assert len(factory.arguments) == 1, "an incomplete run must not restart"


def test_incomplete_status_has_its_own_text(tmp_path: Path) -> None:
    _, _, store, run_id = prepared_root(tmp_path)
    run = store.get_run(run_id)
    model = StatusViewModel.from_run(
        run, worker_alive=False, supervisor_status="incomplete"
    )
    assert "未完成" in model.status_text


def test_unexpected_exit_still_restarts(tmp_path: Path) -> None:
    root, _, store, _ = prepared_root(tmp_path)
    factory = ProcessFactory([FakeProcess(returncode=1), FakeProcess()])
    supervisor = WorkerSupervisor(
        process_factory=factory, store_factory=lambda _workspace: store
    )
    supervisor.start(root, tmp_path / "skill")
    assert supervisor.poll().status == "restarting"


# --- credentials ----------------------------------------------------------

def test_credential_free_environment_strips_provider_keys() -> None:
    environment = credential_free_environment(
        {"GEMINI_API_KEY": "secret", "GOOGLE_API_KEY": "g", "PATH": "x"}
    )
    assert environment == {"PATH": "x"}


def test_installers_do_not_inherit_the_gemini_key(monkeypatch) -> None:
    captured = {}

    def fake_run(args, **kwargs):
        captured.update(kwargs)

        class Result:
            returncode = 0
            stdout = stderr = ""

        return Result()

    monkeypatch.setenv("GEMINI_API_KEY", "secret-value")
    monkeypatch.setattr(setup_environment.subprocess, "run", fake_run)
    setup_environment.run_external(["npm", "install"])

    assert "GEMINI_API_KEY" not in captured["env"]


# --- excel ----------------------------------------------------------------

def test_try_write_excel_reports_locked_workbook(tmp_path: Path) -> None:
    target = tmp_path / "catalog.xlsx"
    target.write_bytes(b"placeholder")

    def locked(_records, _path):
        raise PermissionError("in use")

    assert try_write_excel([], target, writer=locked) is False


def test_try_write_excel_raises_when_first_workbook_cannot_be_created(tmp_path: Path) -> None:
    def locked(_records, _path):
        raise PermissionError("denied")

    with pytest.raises(PermissionError):
        try_write_excel([], tmp_path / "missing.xlsx", writer=locked)


def test_excel_checkpoint_throttles_by_time() -> None:
    now = [0.0]
    writes = []
    checkpoint = ExcelCheckpoint(
        lambda: writes.append(now[0]), every_items=100,
        min_interval_seconds=60, clock=lambda: now[0],
    )
    for count in range(1, 401):
        if count == 300:
            now[0] = 61.0
        checkpoint(count)
    assert writes == [0.0, 61.0]


def test_hyperlinks_are_capped(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(excel_catalog, "MAX_HYPERLINKS", 2)
    for number in range(4):
        (tmp_path / f"clip-{number}.mp4").write_bytes(b"v%d" % number)
    workspace = bootstrap_workspace(tmp_path).workspace
    book = load_workbook(workspace.excel_path)
    try:
        sheet = book["媒體清冊"]
        links = [sheet.cell(row, 3).hyperlink for row in range(2, 6)]
    finally:
        book.close()
    assert sum(link is not None for link in links) == 2


def test_open_excel_does_not_fail_catalog_refresh(tmp_path: Path, monkeypatch) -> None:
    from media_catalog import bootstrap

    (tmp_path / "clip.mp4").write_bytes(b"video")
    bootstrap_workspace(tmp_path)

    def locked(_records, _path):
        raise PermissionError("Excel has the file open")

    monkeypatch.setattr(bootstrap, "write_excel", locked)
    result = bootstrap_workspace(tmp_path)
    assert result.excel_sync_pending is True


def test_stale_temporary_workbooks_are_swept(tmp_path: Path) -> None:
    target = tmp_path / "catalog.xlsx"
    leftover = tmp_path / ".catalog.deadbeef.tmp.xlsx"
    leftover.write_bytes(b"partial")
    old = leftover.stat().st_mtime - 7200
    os.utime(leftover, (old, old))

    write_excel([], target)

    assert not leftover.exists()
    assert target.is_file()


# --- scanning -------------------------------------------------------------

def test_rescan_skips_hashing_unchanged_files(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "clip.mp4").write_bytes(b"video")
    bootstrap_workspace(tmp_path)
    calls = []
    real = scanner_module.sha256_file
    monkeypatch.setattr(
        scanner_module, "sha256_file", lambda path, *a: calls.append(path) or real(path, *a)
    )

    result = bootstrap_workspace(tmp_path)

    assert calls == []
    assert result.scan.existing == 1


def test_rescan_rehashes_modified_file_and_retires_old_row(tmp_path: Path) -> None:
    clip = tmp_path / "clip.mp4"
    clip.write_bytes(b"video")
    bootstrap_workspace(tmp_path)
    clip.write_bytes(b"edited video")
    stat = clip.stat()
    os.utime(clip, ns=(stat.st_atime_ns, stat.st_mtime_ns + 10_000_000))

    result = bootstrap_workspace(tmp_path)
    database = CatalogDatabase(result.workspace.database_path)

    assert result.scan.discovered == 1
    assert result.scan.retired == 1
    assert len(database.list_records()) == 1
    statuses = sorted(r.status for r in database.list_records(include_missing=True))
    assert statuses == sorted([Status.PENDING, Status.MISSING])


def test_deleted_file_is_retired_and_not_requeued(tmp_path: Path) -> None:
    clip = tmp_path / "clip.mp4"
    clip.write_bytes(b"video")
    workspace = bootstrap_workspace(tmp_path).workspace
    clip.unlink()

    result = bootstrap_workspace(tmp_path)
    database = CatalogDatabase(workspace.database_path)

    assert result.scan.retired == 1
    assert database.list_records() == []
    assert database.requeue_failed() == 0


def test_folder_case_rename_does_not_duplicate_rows(tmp_path: Path) -> None:
    root = tmp_path / "media"
    (root / "Trip").mkdir(parents=True)
    (root / "Trip" / "clip.mp4").write_bytes(b"video")
    workspace = bootstrap_workspace(root).workspace
    (root / "Trip").rename(root / "trip")

    bootstrap_workspace(root)
    records = CatalogDatabase(workspace.database_path).list_records()

    assert len(records) == 1
    assert records[0].path.parent.name == "trip"


def test_moved_file_reuses_existing_analysis(tmp_path: Path) -> None:
    (tmp_path / "a").mkdir()
    clip = tmp_path / "a" / "clip.mp4"
    clip.write_bytes(b"video")
    workspace = bootstrap_workspace(tmp_path).workspace
    database = CatalogDatabase(workspace.database_path)
    (record,) = database.list_records()
    database.save_analysis(
        record.id, description="海邊", highlights=("日落",), keywords=("海",)
    )
    (tmp_path / "b").mkdir()
    clip.rename(tmp_path / "b" / "clip.mp4")

    bootstrap_workspace(tmp_path)
    (moved,) = database.list_records()

    assert moved.path.parent.name == "b"
    assert moved.status is Status.ANALYZED
    assert moved.description == "海邊"


def test_unreadable_directory_does_not_abort_scan(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "ok.mp4").write_bytes(b"video")
    (tmp_path / "locked").mkdir()
    real_scandir = os.scandir

    def scandir(path):
        if Path(path).name == "locked":
            raise PermissionError("denied")
        return real_scandir(path)

    monkeypatch.setattr(scanner_module.os, "scandir", scandir)
    result = bootstrap_workspace(tmp_path)

    assert result.scan.supported == 1
    assert result.scan.unreadable == 1


def test_file_locked_during_hash_is_skipped(tmp_path: Path, monkeypatch) -> None:
    (tmp_path / "ok.mp4").write_bytes(b"video")
    (tmp_path / "busy.mp4").write_bytes(b"busy")
    real = scanner_module.sha256_file

    def hash_file(path, *args):
        if Path(path).name == "busy.mp4":
            raise PermissionError("in use")
        return real(path, *args)

    monkeypatch.setattr(scanner_module, "sha256_file", hash_file)
    result = bootstrap_workspace(tmp_path)

    assert result.scan.supported == 1
    assert result.scan.unreadable == 1


def test_media_summary_uses_stored_sizes(tmp_path: Path) -> None:
    (tmp_path / "clip.mp4").write_bytes(b"12345")
    (tmp_path / "photo.jpg").write_bytes(b"123")
    workspace = bootstrap_workspace(tmp_path).workspace
    database = CatalogDatabase(workspace.database_path)

    assert database.media_summary() == (1, 1, 8)
    assert database.media_summary(video_only=True) == (1, 0, 5)


# --- sqlite ---------------------------------------------------------------

def test_connections_are_closed_after_with_block(tmp_path: Path) -> None:
    database = CatalogDatabase(tmp_path / "catalog.sqlite")
    with database._connect() as connection:
        connection.execute("SELECT 1")
    with pytest.raises(sqlite3.ProgrammingError):
        connection.execute("SELECT 1")


def test_run_state_connections_are_closed(tmp_path: Path) -> None:
    store = RunStateStore(tmp_path / "catalog.sqlite")
    with store._connect() as connection:
        connection.execute("SELECT 1")
    with pytest.raises(sqlite3.ProgrammingError):
        connection.execute("SELECT 1")


# --- ui -------------------------------------------------------------------

def test_refresh_loop_survives_database_errors() -> None:
    from types import SimpleNamespace

    from media_catalog.status_ui import StatusApplication

    scheduled = []
    messages = []

    def locked_poll():
        raise sqlite3.OperationalError("database is locked")

    stub = SimpleNamespace(
        supervisor=SimpleNamespace(poll=locked_poll),
        status_var=SimpleNamespace(set=messages.append),
        light=SimpleNamespace(itemconfigure=lambda *a, **k: None),
        light_dot=1,
        root=SimpleNamespace(after=lambda delay, callback: scheduled.append(delay)),
        _refresh=None,
    )

    StatusApplication._refresh(stub)

    assert scheduled == [1000], "the refresh loop must keep running"
    assert "database is locked" in messages[0]
