# -*- coding: utf-8 -*-
"""
影像圖書館 後端服務
- 匯入 Video Library Organizer 產生的「媒體清冊」Excel
- 以 SQLite 保存目錄、挑選、評等、備註、標記點
- 支援 HTTP Range 的影片串流（可拖曳進度）、縮圖、預覽格、轉檔代理檔
啟動：python server.py [--port 8765] [--import 清冊.xlsx ...]
"""
import argparse
import io
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
import openpyxl

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
CACHE = DATA / "cache"
THUMBS = CACHE / "thumbs"
SPRITES = CACHE / "sprites"
PROXIES = CACHE / "proxies"
DB_PATH = DATA / "library.db"
CONFIG_PATH = BASE / "config.json"
for d in (DATA, THUMBS, SPRITES, PROXIES):
    d.mkdir(parents=True, exist_ok=True)

SPRITE_FRAMES = 12
SPRITE_W, SPRITE_H = 160, 90

# Excel 欄位 → 資料庫欄位（未列出的欄位會保留於 extra）
COLUMN_MAP = {
    "狀態": "status",
    "檔名": "filename",
    "完整路徑": "full_path",
    "媒體類型": "media_type",
    "內容描述": "description",
    "重點": "highlights",
    "關鍵字": "keywords",
    "拍攝時間": "shot_time",
    "處理時間": "processed_time",
    "Markdown 路徑": "md_path",
    "備份路徑": "backup_path",
    "錯誤原因": "error",
}
FIELD_LABEL = {v: k for k, v in COLUMN_MAP.items()}
# 面板上的使用者欄位；可選擇是否一併寫回 Excel（會新增欄位）
USER_COLUMNS = {"挑選": "picked", "評等": "rating", "備註": "note"}
USER_LABEL = {v: k for k, v in USER_COLUMNS.items()}
KW_SPLIT = re.compile(r"[、，,;；|\n]+")
HL_SPLIT = re.compile(r"[；;\n]+")

# ---------------------------------------------------------------- config

DEFAULT_CONFIG = {
    "path_maps": [],          # [{"from": "\\\\192.168.1.10\\", "to": "Z:\\"}]
    "ffmpeg": "ffmpeg",
    "ffprobe": "ffprobe",
    "proxy_height": 720,
    "sync_excel": True,         # 修改後即時寫回原始 Excel
    "sync_user_fields": False,  # 挑選／評等／備註也寫回 Excel
    "sync_via_excel_app": True, # 檔案正在 Excel 中開啟時，透過 Excel 直接寫入
}


def load_config():
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        try:
            cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
        except Exception as e:
            print("config.json 讀取失敗：", e)
    return cfg


CONFIG = load_config()


def save_config():
    CONFIG_PATH.write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2), encoding="utf-8")


def _norm(p):
    return (p or "").replace("/", "\\").rstrip("\\").lower()


def resolve_path(full_path):
    """依路徑對應規則換算成本機可讀取的實際路徑"""
    if not full_path:
        return None
    candidates = []
    src = full_path.replace("/", "\\")
    for m in CONFIG.get("path_maps", []):
        f = (m.get("from") or "").replace("/", "\\")
        t = m.get("to") or ""
        if f and src.lower().startswith(f.lower()):
            rest = src[len(f):].lstrip("\\")
            candidates.append(os.path.join(t, *rest.split("\\")) if rest else t)
    candidates.append(full_path)
    for c in candidates:
        try:
            if os.path.isfile(c):
                return c
        except OSError:
            pass
    return None


# ---------------------------------------------------------------- database

_db_lock = threading.RLock()


def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as c:
        c.executescript(
            """
            CREATE TABLE IF NOT EXISTS batches(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              source_name TEXT, sheet TEXT, imported_at TEXT, row_count INTEGER,
              columns TEXT);
            CREATE TABLE IF NOT EXISTS media(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              batch_id INTEGER, row_no INTEGER,
              status TEXT, filename TEXT, full_path TEXT UNIQUE, media_type TEXT,
              description TEXT, highlights TEXT, keywords TEXT,
              shot_time TEXT, processed_time TEXT, md_path TEXT, backup_path TEXT, error TEXT,
              extra TEXT DEFAULT '{}',
              picked INTEGER DEFAULT 0, rating INTEGER DEFAULT 0, note TEXT DEFAULT '',
              duration REAL, width INTEGER, height INTEGER, fps REAL,
              vcodec TEXT, vprofile TEXT, pix_fmt TEXT, acodec TEXT, bitrate INTEGER,
              size INTEGER, creation_time TEXT, probed_at TEXT,
              updated_at TEXT);
            CREATE TABLE IF NOT EXISTS markers(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              media_id INTEGER, t REAL, t_out REAL, label TEXT, created_at TEXT);
            CREATE INDEX IF NOT EXISTS idx_markers_media ON markers(media_id);
            CREATE TABLE IF NOT EXISTS edits(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              media_id INTEGER, field TEXT, old TEXT, new TEXT, at TEXT, source TEXT);
            CREATE INDEX IF NOT EXISTS idx_edits_media ON edits(media_id);
            CREATE TABLE IF NOT EXISTS sync_queue(
              media_id INTEGER, field TEXT, seq INTEGER, PRIMARY KEY(media_id, field));
            """
        )
        # 舊版資料庫升級
        have = {r["name"] for r in c.execute("PRAGMA table_info(batches)")}
        for col, typ in (("source_path", "TEXT"), ("src_mtime", "REAL"), ("last_sync", "TEXT")):
            if col not in have:
                c.execute(f"ALTER TABLE batches ADD COLUMN {col} {typ}")
        have = {r["name"] for r in c.execute("PRAGMA table_info(media)")}
        if "xl_key" not in have:
            c.execute("ALTER TABLE media ADD COLUMN xl_key TEXT")
            c.execute("UPDATE media SET xl_key=full_path")


def now_iso():
    return datetime.now().isoformat(timespec="seconds")


def cell_str(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def user_val(key, v):
    if key == "picked":
        return 0 if v in (None, "", 0) or str(v).strip().lower() in ("0", "否", "no", "false", "n") else 1
    if key == "rating":
        try:
            return max(0, min(5, int(float(v)))) if v not in (None, "") else 0
        except (TypeError, ValueError):
            return 0
    return "" if v is None else str(v)


def import_workbook(src, source_name, source_path=None, only_batch=None):
    """
    src: 檔案路徑或 file-like。source_path 為原始 Excel 在磁碟上的位置（有值才能寫回）。
    以「完整路徑」合併；同一份清冊重新匯入時沿用原本的批次，面板上的挑選／評等／備註／標記點保留。
    """
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    try:
        return _import_wb(wb, source_name, source_path, only_batch)
    finally:
        wb.close()  # read_only 模式會一直佔用檔案，必須關閉，否則 Excel 無法存檔


def _import_wb(wb, source_name, source_path, only_batch):
    total, inserted, updated, changed = 0, 0, 0, 0
    batch_ids = []
    fields = list(COLUMN_MAP.values())
    with _db_lock, db() as c:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":  # 例如 Organizer 的隱藏比對表
                continue
            rows = ws.iter_rows(values_only=True)
            header, header_row = None, 0
            for r in rows:  # 找到含「完整路徑」或「檔名」的標題列
                header_row += 1
                if r and any(cell_str(x) in ("完整路徑", "檔名") for x in r):
                    header = [cell_str(x) or f"欄{i+1}" for i, x in enumerate(r)]
                    break
            if not header:
                continue
            if only_batch:
                b = c.execute("SELECT * FROM batches WHERE id=?", (only_batch,)).fetchone()
                if not b or b["sheet"] != ws.title:
                    continue
            elif source_path:
                b = c.execute("SELECT * FROM batches WHERE source_path=? AND sheet=?",
                              (source_path, ws.title)).fetchone()
            else:
                b = c.execute("SELECT * FROM batches WHERE source_path IS NULL AND source_name=? AND sheet=?",
                              (source_name, ws.title)).fetchone()
            cols_json = json.dumps(header, ensure_ascii=False)
            if b:
                bid = b["id"]
                c.execute("UPDATE batches SET source_name=?, imported_at=?, columns=?, source_path=COALESCE(?, source_path) "
                          "WHERE id=?", (source_name, now_iso(), cols_json, source_path, bid))
            else:
                bid = c.execute(
                    "INSERT INTO batches(source_name, sheet, imported_at, row_count, columns, source_path) "
                    "VALUES(?,?,?,?,?,?)", (source_name, ws.title, now_iso(), 0, cols_json, source_path)).lastrowid
            batch_ids.append(bid)
            n = 0
            for row_no, r in enumerate(rows, start=header_row + 1):
                if not r or all(x is None for x in r):
                    continue
                rec, extra, user = {}, {}, {}
                for h, v in zip(header, r):
                    key = COLUMN_MAP.get(h)
                    if key:
                        rec[key] = cell_str(v)
                    elif h in USER_COLUMNS:
                        user[USER_COLUMNS[h]] = user_val(USER_COLUMNS[h], cell_str(v))
                    elif cell_str(v) is not None:
                        extra[h] = cell_str(v)
                if not rec.get("full_path"):
                    if not rec.get("filename"):
                        continue
                    rec["full_path"] = f"[{source_name}]/{rec['filename']}"
                if not rec.get("filename"):
                    rec["filename"] = re.split(r"[\\/]", rec["full_path"])[-1]
                n += 1
                vals = [rec.get(f) for f in fields]
                exist = c.execute("SELECT * FROM media WHERE full_path=? OR xl_key=? LIMIT 1",
                                  (rec["full_path"], rec["full_path"])).fetchone()
                if exist:
                    # 記錄 Excel 端的變更
                    old_extra = json.loads(exist["extra"] or "{}")
                    diffs = [(f, exist[f], rec.get(f)) for f in fields if (exist[f] or None) != (rec.get(f) or None)]
                    for k in set(old_extra) | set(extra):
                        if (old_extra.get(k) or None) != (extra.get(k) or None):
                            diffs.append(("x:" + k, old_extra.get(k), extra.get(k)))
                    for k, v in user.items():
                        if exist[k] != v:
                            diffs.append((k, exist[k], v))
                    for f, o, nw in diffs:
                        c.execute("INSERT INTO edits(media_id, field, old, new, at, source) VALUES(?,?,?,?,?,?)",
                                  (exist["id"], f, None if o is None else str(o), None if nw is None else str(nw),
                                   now_iso(), "excel"))
                    changed += bool(diffs)
                    sets = ",".join(f"{f}=?" for f in fields)
                    uset = "".join(f", {k}=?" for k in user)
                    c.execute(
                        f"UPDATE media SET {sets}, extra=?, batch_id=?, row_no=?, xl_key=?, updated_at=?{uset} WHERE id=?",
                        vals + [json.dumps(extra, ensure_ascii=False), bid, row_no, rec["full_path"], now_iso()]
                        + list(user.values()) + [exist["id"]],
                    )
                    updated += 1
                else:
                    ucols = "".join(f", {k}" for k in user)
                    c.execute(
                        f"INSERT INTO media({','.join(fields)}, extra, batch_id, row_no, xl_key, updated_at{ucols}) "
                        f"VALUES({','.join('?' * (len(fields) + 5 + len(user)))})",
                        vals + [json.dumps(extra, ensure_ascii=False), bid, row_no, rec["full_path"], now_iso()]
                        + list(user.values()),
                    )
                    inserted += 1
            mtime = os.path.getmtime(source_path) if source_path and os.path.exists(source_path) else None
            c.execute("UPDATE batches SET row_count=?, src_mtime=COALESCE(?, src_mtime) WHERE id=?", (n, mtime, bid))
            total += n
    start_probe_all()
    return {"rows": total, "inserted": inserted, "updated": updated, "changed": changed, "batches": batch_ids,
            "source_path": source_path}


def locate_source(data, filename):
    """瀏覽器上傳無法得知原始位置：在桌面／下載／文件中找同名且內容相同的檔案"""
    import hashlib
    digest = hashlib.sha1(data).hexdigest()
    home = Path.home()
    roots = [home / "Desktop", home / "Downloads", home / "Documents", home / "OneDrive" / "Desktop",
             home / "OneDrive" / "文件", home / "OneDrive" / "Documents", BASE]
    found = []
    for root in roots:
        if not root.is_dir():
            continue
        for p in [root / filename] + list(root.glob(f"*/{filename}")):
            try:
                if p.is_file() and p.stat().st_size == len(data) and \
                        hashlib.sha1(p.read_bytes()).hexdigest() == digest:
                    rp = str(p.resolve())
                    if rp not in found:
                        found.append(rp)
            except OSError:
                pass
    return found[0] if len(found) == 1 else None


# ---------------------------------------------------------------- 欄位編輯 + 寫回 Excel

def field_header(field):
    if field in FIELD_LABEL:
        return FIELD_LABEL[field]
    if field in USER_LABEL:
        return USER_LABEL[field]
    if field.startswith("x:"):
        return field[2:]
    return None


def norm_value(field, v):
    if field in ("picked", "rating"):
        return user_val(field, v)
    if v is None:
        return None if field != "note" else ""
    if isinstance(v, list):
        sep = "；" if field == "highlights" else "、"
        v = sep.join(x.strip() for x in v if str(x).strip())
    s = str(v).replace("\r\n", "\n").strip()
    if field == "note":
        return s
    return s or None


def excel_value(field, v):
    if field == "picked":
        return "是" if v else None
    if field == "rating":
        return v or None
    return v


_sync_seq = [int(time.time() * 1000)]
_sync_event = threading.Event()


def set_fields(media_id, changes, source="edit", c=None):
    """更新欄位、記錄修改歷程、排入 Excel 寫回佇列。回傳實際變更的欄位"""
    own = c is None
    if own:
        c = db()
    try:
        with _db_lock:
            m = c.execute("SELECT * FROM media WHERE id=?", (media_id,)).fetchone()
            if not m:
                raise HTTPException(404, "找不到此媒體")
            extra = json.loads(m["extra"] or "{}")
            done = []
            for field, raw in changes.items():
                if field not in FIELD_LABEL and field not in USER_LABEL and not field.startswith("x:"):
                    continue
                if field.startswith("x:") and (not field[2:].strip() or field[2:] in COLUMN_MAP or field[2:] in USER_COLUMNS):
                    raise HTTPException(400, f"欄位名稱無效：{field[2:]}")
                new = norm_value(field, raw)
                old = extra.get(field[2:]) if field.startswith("x:") else m[field]
                if (old if old not in ("",) else None) == (new if new not in ("",) else None):
                    continue
                if field == "full_path":
                    if not new:
                        raise HTTPException(400, "完整路徑不可為空白")
                    dup = c.execute("SELECT id FROM media WHERE full_path=? AND id<>?", (new, media_id)).fetchone()
                    if dup:
                        raise HTTPException(409, "已有其他媒體使用相同的完整路徑")
                if field.startswith("x:"):
                    if new is None:
                        extra.pop(field[2:], None)
                    else:
                        extra[field[2:]] = new
                    c.execute("UPDATE media SET extra=?, updated_at=? WHERE id=?",
                              (json.dumps(extra, ensure_ascii=False), now_iso(), media_id))
                else:
                    c.execute(f"UPDATE media SET {field}=?, updated_at=? WHERE id=?", (new, now_iso(), media_id))
                c.execute("INSERT INTO edits(media_id, field, old, new, at, source) VALUES(?,?,?,?,?,?)",
                          (media_id, field, None if old is None else str(old), None if new is None else str(new),
                           now_iso(), source))
                if field not in USER_LABEL or CONFIG.get("sync_user_fields"):
                    _sync_seq[0] += 1
                    c.execute("INSERT OR REPLACE INTO sync_queue(media_id, field, seq) VALUES(?,?,?)",
                              (media_id, field, _sync_seq[0]))
                done.append(field)
            if own:
                c.commit()
    finally:
        if own:
            c.close()
    if done:
        _sync_event.set()
    return done


_sync_state = {}   # batch_id -> {"state", "msg", "at", "via"}
_backed_up = set()
BACKUPS = DATA / "backups"


def backup_once(bid, path):
    if bid in _backed_up:
        return
    BACKUPS.mkdir(parents=True, exist_ok=True)
    p = Path(path)
    shutil.copy2(p, BACKUPS / f"{p.stem}_{datetime.now():%Y%m%d_%H%M%S}{p.suffix}")
    _backed_up.add(bid)
    # 每份清冊只保留最近 20 份備份
    olds = sorted(BACKUPS.glob(f"{p.stem}_*{p.suffix}"))
    for o in olds[:-20]:
        try:
            o.unlink()
        except OSError:
            pass


def sync_once():
    from excel_sync import FileLocked, write_entries
    if not CONFIG.get("sync_excel", True):
        return
    with db() as c:
        q = c.execute(
            "SELECT q.media_id, q.field, q.seq, m.batch_id FROM sync_queue q JOIN media m ON m.id=q.media_id"
        ).fetchall()
        # 已不存在的媒體
        c.execute("DELETE FROM sync_queue WHERE media_id NOT IN (SELECT id FROM media)")
    by_batch = {}
    for r in q:
        by_batch.setdefault(r["batch_id"], []).append(r)
    for bid, items in by_batch.items():
        with db() as c:
            b = c.execute("SELECT * FROM batches WHERE id=?", (bid,)).fetchone()
        if not b:
            continue
        path = b["source_path"]
        if not path:
            _sync_state[bid] = {"state": "unlinked", "msg": "尚未連結原始 Excel 檔案", "at": now_iso()}
            continue
        if not os.path.exists(path):
            _sync_state[bid] = {"state": "error", "msg": f"找不到原始 Excel：{path}", "at": now_iso()}
            continue
        entries = {}
        with db() as c:
            for it in items:
                m = c.execute("SELECT * FROM media WHERE id=?", (it["media_id"],)).fetchone()
                e = entries.get(m["id"])
                if not e:
                    extra = json.loads(m["extra"] or "{}")
                    full = {h: m[f] for h, f in COLUMN_MAP.items()}
                    full.update(extra)
                    e = entries[m["id"]] = {"media_id": m["id"], "row_no": m["row_no"], "key": m["xl_key"] or m["full_path"],
                                            "cells": {}, "full": full, "_m": m, "_extra": extra}
                f = it["field"]
                if f.startswith("x:"):
                    v = e["_extra"].get(f[2:])
                else:
                    v = excel_value(f, m[f])
                e["cells"][field_header(f)] = v
        try:
            try:
                backup_once(bid, path)
            except PermissionError as ex:  # 連讀取都被擋：先不寫入，等檔案釋放（確保寫入前一定有備份）
                raise FileLocked("無法讀取檔案以建立備份")
            payload = [{k: v for k, v in e.items() if not k.startswith("_")} for e in entries.values()]
            rows, via = write_entries(path, b["sheet"], payload, allow_com=CONFIG.get("sync_via_excel_app", True))
        except FileLocked as ex:
            _sync_state[bid] = {"state": "locked", "msg": "Excel 檔案正被其他程式開啟，關閉後會自動寫入（" + str(ex)[:80] + "）",
                                "at": now_iso()}
            continue
        except Exception as ex:
            _sync_state[bid] = {"state": "error", "msg": f"寫入失敗：{ex}", "at": now_iso()}
            continue
        with _db_lock, db() as c:
            for it in items:
                c.execute("DELETE FROM sync_queue WHERE media_id=? AND field=? AND seq=?",
                          (it["media_id"], it["field"], it["seq"]))
            for mid, r in rows.items():
                c.execute("UPDATE media SET row_no=?, xl_key=full_path WHERE id=?", (r, mid))
            c.execute("UPDATE batches SET last_sync=?, src_mtime=? WHERE id=?",
                      (now_iso(), os.path.getmtime(path), bid))
        _sync_state[bid] = {"state": "ok", "msg": "已同步" + ("（透過 Excel 寫入已開啟的檔案）" if via == "excel" else ""),
                            "at": now_iso(), "via": via}


def sync_worker():
    while True:
        _sync_event.wait(timeout=3)
        _sync_event.clear()
        time.sleep(0.3)  # 合併連續修改
        try:
            sync_once()
        except Exception as e:
            print("同步錯誤：", e)


def sync_status():
    out = []
    with db() as c:
        pend = {r["batch_id"]: r["n"] for r in c.execute(
            "SELECT m.batch_id, COUNT(*) n FROM sync_queue q JOIN media m ON m.id=q.media_id GROUP BY m.batch_id")}
        for b in c.execute("SELECT * FROM batches ORDER BY id"):
            st = dict(_sync_state.get(b["id"]) or {})
            n = pend.get(b["id"], 0)
            path = b["source_path"]
            ext = False
            if path and os.path.exists(path) and b["src_mtime"]:
                ext = os.path.getmtime(path) > b["src_mtime"] + 1
            if not path:
                st = {"state": "unlinked", "msg": "尚未連結原始 Excel 檔案，修改只存在面板中"}
            elif not n and st.get("state") in (None, "locked", "pending"):
                st = {"state": "ok", "msg": "已同步", "at": b["last_sync"]}
            elif n and st.get("state") in (None, "ok"):
                st = {"state": "pending", "msg": "等待寫入"}
            if not CONFIG.get("sync_excel", True) and path:
                st = {"state": "off", "msg": "已關閉自動寫回 Excel"}
            out.append({"id": b["id"], "source_name": b["source_name"], "source_path": path, "pending": n,
                        "last_sync": b["last_sync"], "external_changed": ext and not n, **st})
    return out


# ---------------------------------------------------------------- ffmpeg helpers

def ff(name):
    exe = CONFIG.get(name) or name
    return shutil.which(exe) or exe


NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0


def run(cmd, timeout=120):
    return subprocess.run(cmd, capture_output=True, timeout=timeout, creationflags=NO_WINDOW)


def probe(media_id):
    with db() as c:
        m = c.execute("SELECT id, full_path FROM media WHERE id=?", (media_id,)).fetchone()
    if not m:
        return
    p = resolve_path(m["full_path"])
    if not p:
        return
    try:
        r = run([ff("ffprobe"), "-v", "error", "-print_format", "json", "-show_format", "-show_streams", p], 60)
        info = json.loads(r.stdout.decode("utf-8", "replace") or "{}")
    except Exception as e:
        print("ffprobe 失敗", p, e)
        return
    fmt = info.get("format", {})
    vs = next((s for s in info.get("streams", []) if s.get("codec_type") == "video"), {})
    as_ = next((s for s in info.get("streams", []) if s.get("codec_type") == "audio"), {})
    fps = None
    try:
        a, b = (vs.get("avg_frame_rate") or vs.get("r_frame_rate") or "0/1").split("/")
        fps = round(float(a) / float(b), 3) if float(b) else None
    except Exception:
        pass
    tags = fmt.get("tags", {}) or {}
    with _db_lock, db() as c:
        c.execute(
            """UPDATE media SET duration=?, width=?, height=?, fps=?, vcodec=?, vprofile=?, pix_fmt=?,
               acodec=?, bitrate=?, size=?, creation_time=?, probed_at=? WHERE id=?""",
            (
                float(fmt["duration"]) if fmt.get("duration") else None,
                vs.get("width"), vs.get("height"), fps,
                vs.get("codec_name"), vs.get("profile"), vs.get("pix_fmt"),
                as_.get("codec_name"),
                int(fmt["bit_rate"]) if fmt.get("bit_rate") else None,
                os.path.getsize(p),
                tags.get("creation_time"),
                now_iso(), media_id,
            ),
        )


_probe_thread = None


def start_probe_all(force=False):
    global _probe_thread
    if _probe_thread and _probe_thread.is_alive():
        return

    def work():
        with db() as c:
            q = "SELECT id FROM media" + ("" if force else " WHERE probed_at IS NULL")
            ids = [r["id"] for r in c.execute(q)]
        for i in ids:
            probe(i)

    _probe_thread = threading.Thread(target=work, daemon=True)
    _probe_thread.start()


_locks = {}
_locks_guard = threading.Lock()


def key_lock(key):
    with _locks_guard:
        return _locks.setdefault(key, threading.Lock())


def media_row(media_id):
    with db() as c:
        m = c.execute("SELECT * FROM media WHERE id=?", (media_id,)).fetchone()
    if not m:
        raise HTTPException(404, "找不到此媒體")
    return m


def playable_in_browser(m):
    vc = (m["vcodec"] or "").lower()
    pf = (m["pix_fmt"] or "").lower()
    if not vc:
        return True  # 未知，交給瀏覽器嘗試
    if vc not in ("h264", "vp8", "vp9", "av1", "hevc"):
        return False
    if vc == "h264" and pf and pf not in ("yuv420p", "yuvj420p"):
        return False  # 例如 XAVC 4:2:2 10-bit
    return True


# ---------------------------------------------------------------- proxies

_proxy_jobs = {}  # id -> {"state": running|done|error, "progress": 0-1, "msg": str}


def proxy_path(media_id):
    return PROXIES / f"{media_id}.mp4"


def start_proxy(media_id):
    m = media_row(media_id)
    src = resolve_path(m["full_path"])
    if not src:
        raise HTTPException(404, "原始檔案無法存取")
    job = _proxy_jobs.get(media_id)
    if job and job["state"] == "running":
        return job
    if proxy_path(media_id).exists():
        return {"state": "done", "progress": 1}
    job = {"state": "running", "progress": 0.0, "msg": ""}
    _proxy_jobs[media_id] = job
    dur = m["duration"] or 0

    def work():
        tmp = PROXIES / f"{media_id}.part.mp4"
        h = int(CONFIG.get("proxy_height") or 720)
        cmd = [ff("ffmpeg"), "-y", "-hide_banner", "-i", src,
               "-vf", f"scale=-2:'min({h},ih)'", "-c:v", "libx264", "-preset", "veryfast", "-crf", "23",
               "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart",
               "-progress", "pipe:1", "-nostats", str(tmp)]
        try:
            p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
            for line in p.stdout:
                line = line.decode("utf-8", "replace").strip()
                if line.startswith("out_time_us=") and dur:
                    try:
                        job["progress"] = min(0.99, int(line.split("=")[1]) / 1e6 / dur)
                    except ValueError:
                        pass
            p.wait()
            if p.returncode == 0:
                tmp.replace(proxy_path(media_id))
                job.update(state="done", progress=1)
            else:
                job.update(state="error", msg=f"ffmpeg 結束代碼 {p.returncode}")
        except Exception as e:
            job.update(state="error", msg=str(e))
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    threading.Thread(target=work, daemon=True).start()
    return job


# ---------------------------------------------------------------- app

app = FastAPI(title="影像圖書館")


@app.middleware("http")
async def no_cache_static(request, call_next):
    resp = await call_next(request)
    if not request.url.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-cache"
    return resp


def serialize(m, root_dirs=None):
    d = dict(m)
    d["extra"] = json.loads(d.get("extra") or "{}")
    real = resolve_path(d["full_path"])
    d["available"] = bool(real)
    d["has_proxy"] = proxy_path(d["id"]).exists()
    d["playable"] = playable_in_browser(m)
    return d


def split_dir(p):
    parts = re.split(r"[\\/]+", p or "")
    return parts[:-1]


@app.get("/api/media")
def list_media():
    with db() as c:
        rows = c.execute("SELECT * FROM media ORDER BY full_path").fetchall()
        mk = {}
        for r in c.execute("SELECT media_id, COUNT(*) n FROM markers GROUP BY media_id"):
            mk[r["media_id"]] = r["n"]
        batches = [dict(b) for b in c.execute("SELECT * FROM batches ORDER BY id")]
    items = [serialize(r) for r in rows]
    # 以共同上層目錄為根，算出相對資料夾（供左側資料夾樹使用）
    dirs = [split_dir(i["full_path"]) for i in items]
    common = []
    if dirs:
        for parts in zip(*dirs):
            if all(x.lower() == parts[0].lower() for x in parts):
                common.append(parts[0])
            else:
                break
    if len(items) == 1:
        common = dirs[0][:-1] if dirs[0] else []
    for i, d in zip(items, dirs):
        i["rel_dir"] = "/".join(d[len(common):])
        i["marker_count"] = mk.get(i["id"], 0)
    for b in batches:
        b["columns"] = json.loads(b.get("columns") or "[]")
    root = "\\".join(common)
    if common and common[0] == "":  # UNC 路徑 \\server\share
        root = "\\" + root
    return {"items": items, "root": root, "root_name": common[-1] if common else "", "batches": batches,
            "field_labels": FIELD_LABEL}


@app.get("/api/media/{media_id}")
def get_media(media_id: int):
    m = media_row(media_id)
    d = serialize(m)
    with db() as c:
        d["markers"] = [dict(x) for x in c.execute("SELECT * FROM markers WHERE media_id=? ORDER BY t", (media_id,))]
    d["resolved_path"] = resolve_path(m["full_path"])
    d["proxy_job"] = _proxy_jobs.get(media_id)
    return d


def body_changes(body):
    """接受 {"fields": {...}} 或舊格式 {"status": ...}"""
    ch = dict(body.get("fields") or {})
    for k in ("status", "picked", "rating", "note"):
        if k in body:
            ch[k] = body[k]
    return ch


@app.patch("/api/media/{media_id}")
async def patch_media(media_id: int, request: Request):
    ch = body_changes(await request.json())
    if not ch:
        raise HTTPException(400, "沒有可更新的欄位")
    changed = set_fields(media_id, ch)
    return get_media(media_id) | {"changed": changed}


@app.post("/api/media/bulk")
async def bulk_update(request: Request):
    body = await request.json()
    ids = [int(i) for i in body.get("ids", [])]
    ch = dict(body.get("set") or {})
    if not ids or not ch:
        raise HTTPException(400, "參數不足")
    n = 0
    with db() as c:
        for i in ids:
            n += bool(set_fields(i, ch, "bulk", c))
    _sync_event.set()
    return {"ok": True, "count": n}


def split_list(field, s):
    return [x.strip() for x in (HL_SPLIT if field == "highlights" else KW_SPLIT).split(s or "") if x.strip()]


@app.post("/api/media/bulk_keywords")
async def bulk_keywords(request: Request):
    """對多筆媒體新增／移除關鍵字"""
    b = await request.json()
    ids = [int(i) for i in b.get("ids", [])]
    add = [x for x in (b.get("add") or []) if x.strip()]
    remove = set(x.strip() for x in (b.get("remove") or []))
    field = b.get("field") or "keywords"
    n = 0
    with db() as c:
        for i in ids:
            m = c.execute(f"SELECT {field} FROM media WHERE id=?", (i,)).fetchone()
            if not m:
                continue
            cur = split_list(field, m[field])
            new = [x for x in cur if x not in remove] + [x.strip() for x in add if x.strip() not in cur]
            new = list(dict.fromkeys(new))
            if new != cur:
                n += bool(set_fields(i, {field: new}, "bulk", c))
        c.commit()
    _sync_event.set()
    return {"ok": True, "count": n}


@app.get("/api/keywords")
def list_keywords():
    cnt = {}
    with db() as c:
        for r in c.execute("SELECT keywords FROM media"):
            for k in split_list("keywords", r["keywords"]):
                cnt[k] = cnt.get(k, 0) + 1
    return sorted(([k, v] for k, v in cnt.items()), key=lambda x: -x[1])


@app.post("/api/keywords")
async def keyword_op(request: Request):
    """全域重新命名／合併／刪除關鍵字：{"from": "舊", "to": "新"}；to 為空白表示刪除"""
    b = await request.json()
    src = (b.get("from") or "").strip()
    dst = (b.get("to") or "").strip()
    if not src:
        raise HTTPException(400, "缺少要修改的關鍵字")
    n = 0
    with db() as c:
        for r in c.execute("SELECT id, keywords FROM media").fetchall():
            cur = split_list("keywords", r["keywords"])
            if src not in cur:
                continue
            new = list(dict.fromkeys([dst if x == src else x for x in cur if dst or x != src]))
            n += bool(set_fields(r["id"], {"keywords": new}, "keyword", c))
        c.commit()
    _sync_event.set()
    return {"ok": True, "count": n}


@app.get("/api/media/{media_id}/edits")
def media_edits(media_id: int):
    with db() as c:
        return [dict(r) for r in c.execute("SELECT * FROM edits WHERE media_id=? ORDER BY id DESC LIMIT 300",
                                           (media_id,))]


@app.post("/api/edits/{edit_id}/revert")
def revert_edit(edit_id: int):
    with db() as c:
        e = c.execute("SELECT * FROM edits WHERE id=?", (edit_id,)).fetchone()
    if not e:
        raise HTTPException(404, "找不到此修改紀錄")
    set_fields(e["media_id"], {e["field"]: e["old"]}, "revert")
    return get_media(e["media_id"])



@app.post("/api/media/{media_id}/markers")
async def add_marker(media_id: int, request: Request):
    b = await request.json()
    with _db_lock, db() as c:
        c.execute(
            "INSERT INTO markers(media_id, t, t_out, label, created_at) VALUES(?,?,?,?,?)",
            (media_id, float(b.get("t", 0)), b.get("t_out"), b.get("label") or "", now_iso()),
        )
    return get_media(media_id)


@app.patch("/api/markers/{marker_id}")
async def edit_marker(marker_id: int, request: Request):
    b = await request.json()
    sets = {k: b[k] for k in ("t", "t_out", "label") if k in b}
    with _db_lock, db() as c:
        if sets:
            c.execute(f"UPDATE markers SET {','.join(f'{k}=?' for k in sets)} WHERE id=?",
                      list(sets.values()) + [marker_id])
    return {"ok": True}


@app.delete("/api/markers/{marker_id}")
def del_marker(marker_id: int):
    with _db_lock, db() as c:
        c.execute("DELETE FROM markers WHERE id=?", (marker_id,))
    return {"ok": True}


@app.post("/api/import")
async def api_import(file: UploadFile = File(...)):
    data = await file.read()
    path = locate_source(data, file.filename)
    try:
        res = import_workbook(path or io.BytesIO(data), file.filename, path)
    except Exception as e:
        raise HTTPException(400, f"無法讀取 Excel：{e}")
    return res


def check_xlsx_path(p):
    p = (p or "").strip().strip('"')
    if not p or not os.path.isfile(p):
        raise HTTPException(400, f"找不到檔案：{p}")
    if not p.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "只支援 .xlsx / .xlsm")
    return os.path.abspath(p)


@app.post("/api/import_path")
async def api_import_path(request: Request):
    p = check_xlsx_path((await request.json()).get("path"))
    try:
        return import_workbook(p, os.path.basename(p), p)
    except PermissionError:
        raise HTTPException(423, "檔案被鎖定，無法讀取")
    except Exception as e:
        raise HTTPException(400, f"無法讀取 Excel：{e}")


@app.post("/api/batches/{batch_id}/link")
async def link_batch(batch_id: int, request: Request):
    """為上傳匯入的清冊指定原始 Excel 位置，之後的修改會寫回該檔案"""
    p = check_xlsx_path((await request.json()).get("path"))
    with db() as c:
        b = c.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
    if not b:
        raise HTTPException(404, "找不到此批次")
    wb = openpyxl.load_workbook(p, read_only=True)
    names = wb.sheetnames
    wb.close()
    if b["sheet"] not in names:
        raise HTTPException(400, f"此 Excel 沒有工作表「{b['sheet']}」，請確認是同一份清冊")
    with _db_lock, db() as c:
        c.execute("UPDATE batches SET source_path=?, src_mtime=? WHERE id=?", (p, os.path.getmtime(p), batch_id))
    _sync_state.pop(batch_id, None)
    _sync_event.set()
    return {"ok": True, "source_path": p}


@app.post("/api/batches/{batch_id}/reload")
async def reload_batch(batch_id: int, request: Request):
    """從原始 Excel 重新載入（套用在 Excel 中做的修改）"""
    force = (await request.json()).get("force")
    with db() as c:
        b = c.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        pend = c.execute("SELECT COUNT(*) n FROM sync_queue q JOIN media m ON m.id=q.media_id WHERE m.batch_id=?",
                         (batch_id,)).fetchone()["n"]
    if not b or not b["source_path"]:
        raise HTTPException(400, "此清冊尚未連結原始 Excel")
    if pend and not force:
        raise HTTPException(409, f"還有 {pend} 項修改尚未寫入 Excel，重新載入會覆蓋這些修改")
    if pend:
        with _db_lock, db() as c:
            c.execute("DELETE FROM sync_queue WHERE media_id IN (SELECT id FROM media WHERE batch_id=?)", (batch_id,))
    try:
        return import_workbook(b["source_path"], b["source_name"], b["source_path"], only_batch=batch_id)
    except PermissionError:
        raise HTTPException(423, "檔案被鎖定，無法讀取")


@app.get("/api/sync")
def api_sync():
    return {"batches": sync_status()}


@app.post("/api/sync")
def api_sync_now():
    _sync_event.set()
    return {"ok": True}


@app.delete("/api/batches/{batch_id}")
def delete_batch(batch_id: int):
    with _db_lock, db() as c:
        ids = [r["id"] for r in c.execute("SELECT id FROM media WHERE batch_id=?", (batch_id,))]
        for t in ("markers", "sync_queue", "edits"):
            c.execute(f"DELETE FROM {t} WHERE media_id IN (SELECT id FROM media WHERE batch_id=?)", (batch_id,))
        c.execute("DELETE FROM media WHERE batch_id=?", (batch_id,))
        c.execute("DELETE FROM batches WHERE id=?", (batch_id,))
    return {"ok": True, "removed": len(ids)}



@app.get("/api/config")
def get_config():
    return CONFIG


@app.put("/api/config")
async def put_config(request: Request):
    try:
        b = await request.json()
    except ValueError:
        raise HTTPException(400, "設定格式錯誤")
    if "path_maps" in b:
        CONFIG["path_maps"] = [
            {"from": m.get("from", "").strip(), "to": m.get("to", "").strip()}
            for m in b["path_maps"] if m.get("from", "").strip()
        ]
    for k in ("ffmpeg", "ffprobe", "proxy_height", "sync_excel", "sync_user_fields", "sync_via_excel_app"):
        if k in b:
            CONFIG[k] = b[k]
    save_config()
    _sync_event.set()
    start_probe_all()
    return CONFIG


@app.post("/api/rescan")
def rescan():
    start_probe_all(force=True)
    return {"ok": True}


@app.get("/api/status")
def status():
    ok = {}
    for k in ("ffmpeg", "ffprobe"):
        ok[k] = bool(shutil.which(CONFIG.get(k) or k))
    return {"tools": ok, "probing": bool(_probe_thread and _probe_thread.is_alive())}


# ---- 媒體檔案

def media_file_response(path, request):
    ctype = mimetypes.guess_type(path)[0] or "application/octet-stream"
    if path.lower().endswith((".mov", ".m4v")):
        ctype = "video/mp4"
    return FileResponse(path, media_type=ctype, headers={"Cache-Control": "no-cache"})


@app.get("/api/stream/{media_id}")
def stream(media_id: int, request: Request, proxy: int = 0):
    m = media_row(media_id)
    if proxy:
        p = proxy_path(media_id)
        if not p.exists():
            raise HTTPException(404, "代理檔尚未產生")
        return media_file_response(str(p), request)
    p = resolve_path(m["full_path"])
    if not p:
        raise HTTPException(404, "原始檔案無法存取，請檢查路徑對應設定")
    return media_file_response(p, request)


@app.get("/api/download/{media_id}")
def download(media_id: int):
    m = media_row(media_id)
    p = resolve_path(m["full_path"])
    if not p:
        raise HTTPException(404, "原始檔案無法存取")
    return FileResponse(p, filename=m["filename"])


@app.post("/api/proxy/{media_id}")
def make_proxy(media_id: int):
    return start_proxy(media_id)


@app.get("/api/proxy/{media_id}")
def proxy_status(media_id: int):
    if proxy_path(media_id).exists():
        return {"state": "done", "progress": 1}
    return _proxy_jobs.get(media_id) or {"state": "none"}


def seek_points(dur, n):
    if not dur or dur <= 0:
        return [0] * n
    return [dur * (i + 0.5) / n for i in range(n)]


@app.get("/api/thumb/{media_id}")
def thumb(media_id: int):
    out = THUMBS / f"{media_id}.jpg"
    if not out.exists():
        m = media_row(media_id)
        src = resolve_path(m["full_path"])
        if not src:
            raise HTTPException(404, "檔案無法存取")
        with key_lock(("thumb", media_id)):
            if not out.exists():
                if not m["probed_at"]:
                    probe(media_id)
                    m = media_row(media_id)
                t = min(max((m["duration"] or 0) * 0.15, 0), 30)
                tmp = THUMBS / f"{media_id}.part.jpg"
                r = run([ff("ffmpeg"), "-y", "-hide_banner", "-loglevel", "error", "-ss", f"{t:.2f}", "-i", src,
                         "-frames:v", "1", "-vf", "scale=480:-2", "-q:v", "4", str(tmp)], 60)
                if r.returncode != 0 or not tmp.exists():
                    raise HTTPException(500, "縮圖產生失敗")
                tmp.replace(out)
    return FileResponse(out, media_type="image/jpeg", headers={"Cache-Control": "max-age=86400"})


@app.get("/api/sprite/{media_id}")
def sprite(media_id: int):
    """水平排列的預覽格（滑鼠移動預覽 / 進度條懸停畫面）"""
    out = SPRITES / f"{media_id}.jpg"
    if not out.exists():
        m = media_row(media_id)
        src = resolve_path(m["full_path"])
        if not src:
            raise HTTPException(404, "檔案無法存取")
        with key_lock(("sprite", media_id)):
            if not out.exists():
                if not m["probed_at"]:
                    probe(media_id)
                    m = media_row(media_id)
                pts = seek_points(m["duration"], SPRITE_FRAMES)
                cmd = [ff("ffmpeg"), "-y", "-hide_banner", "-loglevel", "error"]
                for t in pts:
                    cmd += ["-ss", f"{t:.2f}", "-i", src]
                chains = []
                for i in range(len(pts)):
                    chains.append(
                        f"[{i}:v]trim=end_frame=1,scale={SPRITE_W}:{SPRITE_H}:force_original_aspect_ratio=decrease,"
                        f"pad={SPRITE_W}:{SPRITE_H}:(ow-iw)/2:(oh-ih)/2,setsar=1[v{i}]"
                    )
                chains.append("".join(f"[v{i}]" for i in range(len(pts))) + f"hstack=inputs={len(pts)}[out]")
                tmp = SPRITES / f"{media_id}.part.jpg"
                cmd += ["-filter_complex", ";".join(chains), "-map", "[out]", "-frames:v", "1", "-q:v", "5", str(tmp)]
                r = run(cmd, 180)
                if r.returncode != 0 or not tmp.exists():
                    print(r.stderr.decode("utf-8", "replace")[-500:])
                    raise HTTPException(500, "預覽格產生失敗")
                tmp.replace(out)
    return FileResponse(out, media_type="image/jpeg",
                        headers={"Cache-Control": "max-age=86400", "X-Frames": str(SPRITE_FRAMES)})


@app.get("/api/markdown/{media_id}")
def markdown(media_id: int):
    m = media_row(media_id)
    if not m["md_path"]:
        raise HTTPException(404, "此筆沒有 Markdown 路徑")
    p = resolve_path(m["md_path"])
    if not p:
        raise HTTPException(404, "Markdown 檔案無法存取")
    return Response(Path(p).read_text(encoding="utf-8", errors="replace"), media_type="text/plain; charset=utf-8")


def fmt_tc(t):
    if t is None:
        return ""
    t = float(t)
    return f"{int(t // 3600):02d}:{int(t % 3600 // 60):02d}:{t % 60:06.3f}"


@app.get("/api/export")
def export(ids: str = "", picked: int = 0):
    """匯出清冊（含挑選、評等、備註、標記點、技術資訊）"""
    with db() as c:
        q = "SELECT * FROM media"
        args = []
        if ids:
            id_list = [int(x) for x in ids.split(",") if x.strip().isdigit()]
            q += f" WHERE id IN ({','.join('?' * len(id_list))})"
            args = id_list
        elif picked:
            q += " WHERE picked=1"
        rows = c.execute(q + " ORDER BY full_path", args).fetchall()
        markers = {}
        for mk in c.execute("SELECT * FROM markers ORDER BY t"):
            markers.setdefault(mk["media_id"], []).append(mk)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "媒體清冊"
    extra_keys = []
    for r in rows:
        for k in json.loads(r["extra"] or "{}"):
            if k not in extra_keys:
                extra_keys.append(k)
    head = list(COLUMN_MAP.keys()) + extra_keys + ["挑選", "評等", "備註", "標記點", "時長", "解析度", "影格率",
                                                 "視訊編碼", "音訊編碼", "檔案大小(MB)"]
    ws.append(head)
    for r in rows:
        ex = json.loads(r["extra"] or "{}")
        mks = "\n".join(
            f"{fmt_tc(x['t'])}{(' → ' + fmt_tc(x['t_out'])) if x['t_out'] is not None else ''}  {x['label'] or ''}"
            for x in markers.get(r["id"], [])
        )
        ws.append([r[f] for f in COLUMN_MAP.values()] + [ex.get(k) for k in extra_keys] + [
            "是" if r["picked"] else "", r["rating"] or "", r["note"] or "", mks,
            fmt_tc(r["duration"]) if r["duration"] else "",
            f"{r['width']}x{r['height']}" if r["width"] else "", r["fps"] or "",
            r["vcodec"] or "", r["acodec"] or "",
            round(r["size"] / 1048576, 1) if r["size"] else "",
        ])
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUV", [9, 16, 60, 12, 60, 40, 40, 20, 26, 30, 30, 20]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"媒體清冊_匯出_{datetime.now():%Y%m%d_%H%M}.xlsx"
    from urllib.parse import quote
    return Response(buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


app.mount("/", StaticFiles(directory=BASE / "static", html=True), name="static")


def main():
    ap = argparse.ArgumentParser(description="影像圖書館 後端")
    ap.add_argument("--host", default="127.0.0.1", help="要讓區網其他電腦使用請設為 0.0.0.0")
    ap.add_argument("--port", type=int, default=8765)
    ap.add_argument("--import", dest="imports", nargs="*", default=[], help="啟動時匯入的 Excel 清冊")
    ap.add_argument("--open", action="store_true", help="啟動後自動開啟瀏覽器")
    a = ap.parse_args()
    init_db()
    for f in a.imports:
        p = os.path.abspath(f)
        print("匯入", p, import_workbook(p, os.path.basename(p), p))
    start_probe_all()
    threading.Thread(target=sync_worker, daemon=True).start()
    if a.open:
        import webbrowser
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{a.port}/")).start()
    import uvicorn
    print(f"影像圖書館已啟動：http://{a.host}:{a.port}/")
    uvicorn.run(app, host=a.host, port=a.port, log_level="warning")


init_db()

if __name__ == "__main__":
    main()
