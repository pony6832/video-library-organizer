# -*- coding: utf-8 -*-
"""
將面板上的修改寫回原始 Excel 清冊。
- 只寫入被修改的儲存格，其他內容（含使用者在 Excel 中做的修改、格式）保持不變
- 檔案未被開啟：以 openpyxl 寫入暫存檔後置換
- 檔案正被 Microsoft Excel 開啟：透過 COM 直接寫入那份已開啟的活頁簿並儲存（不會自行開啟檔案）
"""
import os
from datetime import datetime

import openpyxl

KEY_HEADERS = ("完整路徑", "檔名")


class FileLocked(Exception):
    """檔案被其他程式（通常是 Excel）鎖定，稍後重試"""


def _s(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def _norm_key(v):
    s = _s(v)
    return s.replace("/", "\\").lower() if s else None


class _Sheet:
    """儲存格存取介面，座標為 1 起算"""

    def get(self, r, c): raise NotImplementedError
    def set(self, r, c, v): raise NotImplementedError
    max_row = 0
    max_col = 0


class OpenpyxlSheet(_Sheet):
    def __init__(self, ws):
        self.ws = ws

    @property
    def max_row(self):
        return self.ws.max_row

    @property
    def max_col(self):
        return self.ws.max_column

    def get(self, r, c):
        return self.ws.cell(row=r, column=c).value

    def set(self, r, c, v):
        self.ws.cell(row=r, column=c).value = v


class ComSheet(_Sheet):
    def __init__(self, ws):
        self.ws = ws
        used = ws.UsedRange
        self.r0, self.c0 = used.Row, used.Column
        vals = used.Value
        if not isinstance(vals, tuple):
            vals = ((vals,),)
        self.vals = [list(r) for r in vals]
        self.max_row = self.r0 + len(self.vals) - 1
        self.max_col = self.c0 + (len(self.vals[0]) if self.vals else 1) - 1

    def get(self, r, c):
        i, j = r - self.r0, c - self.c0
        if 0 <= i < len(self.vals) and 0 <= j < len(self.vals[i]):
            v = self.vals[i][j]
            # COM 回傳的日期是 pywintypes.datetime
            return v.isoformat() if hasattr(v, "isoformat") and not isinstance(v, str) else v
        return None

    def set(self, r, c, v):
        self.ws.Cells(r, c).Value = "" if v is None else v
        self.max_row = max(self.max_row, r)
        self.max_col = max(self.max_col, c)


def find_header(sheet, limit=30):
    for r in range(1, min(sheet.max_row, limit) + 1):
        row = [_s(sheet.get(r, c)) for c in range(1, sheet.max_col + 1)]
        if any(x in KEY_HEADERS for x in row):
            return r, {h: i + 1 for i, h in enumerate(row) if h}
    raise ValueError("Excel 中找不到含「完整路徑」或「檔名」的標題列")


def apply_entries(sheet, entries):
    """
    entries: [{media_id, row_no, key, cells: {標題: 值}, full: {標題: 值}}]
      key  = 目前 Excel 中「完整路徑」的值（用來找到正確的列）
      full = 找不到該列時，新增一列所需的完整資料
    回傳 {media_id: 實際列號}
    """
    hr, cols = find_header(sheet)
    key_col = cols.get("完整路徑") or cols.get("檔名")
    key_header = "完整路徑" if "完整路徑" in cols else "檔名"

    def col_for(h):
        if h not in cols:
            c = max([sheet.max_col] + list(cols.values())) + 1
            sheet.set(hr, c, h)
            cols[h] = c
        return cols[h]

    index = None
    rows = {}
    for e in entries:
        key = _norm_key(e["key"])
        r = e.get("row_no")
        if not (r and r > hr and _norm_key(sheet.get(r, key_col)) == key):
            if index is None:
                index = {}
                for rr in range(hr + 1, sheet.max_row + 1):
                    k = _norm_key(sheet.get(rr, key_col))
                    if k and k not in index:
                        index[k] = rr
            r = index.get(key)
        if not r:
            # Excel 中已不存在此列（被刪除或改過路徑）→ 新增到最後
            r = max(sheet.max_row, hr) + 1
            for h, v in e["full"].items():
                if v is not None:
                    sheet.set(r, col_for(h), v)
            if index is not None:
                index[_norm_key(e["full"].get(key_header))] = r
        for h, v in e["cells"].items():
            sheet.set(r, col_for(h), v)
        rows[e["media_id"]] = r
    return rows


def write_with_openpyxl(path, sheet_name, entries):
    keep_vba = path.lower().endswith(".xlsm")
    try:
        wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    except PermissionError as e:
        raise FileLocked(str(e))
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    rows = apply_entries(OpenpyxlSheet(ws), entries)
    tmp = path + ".vltmp"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as e:
        raise FileLocked(str(e))
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
    return rows


def write_with_excel_com(path, sheet_name, entries):
    """僅在檔案已於 Excel 中開啟時使用；找不到已開啟的活頁簿則丟出 FileLocked"""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        raise FileLocked("檔案被鎖定，且未安裝 pywin32")
    pythoncom.CoInitialize()
    try:
        try:
            xl = win32com.client.GetActiveObject("Excel.Application")
        except Exception:
            raise FileLocked("檔案被其他程式鎖定")
        target = None
        want = os.path.normcase(os.path.abspath(path))
        for wb in xl.Workbooks:
            try:
                if os.path.normcase(os.path.abspath(wb.FullName)) == want:
                    target = wb
                    break
            except Exception:
                continue
        if target is None:
            raise FileLocked("檔案被其他程式鎖定（不是由 Excel 開啟）")
        if target.ReadOnly:
            raise FileLocked("Excel 以唯讀模式開啟此檔案")
        try:
            ws = target.Worksheets(sheet_name)
        except Exception:
            ws = target.Worksheets(1)
        rows = apply_entries(ComSheet(ws), entries)
        target.Save()
        return rows
    finally:
        pythoncom.CoUninitialize()


def write_entries(path, sheet_name, entries, allow_com=True):
    """回傳 (列號對照, 寫入方式)"""
    try:
        return write_with_openpyxl(path, sheet_name, entries), "file"
    except FileLocked:
        if not allow_com:
            raise
        return write_with_excel_com(path, sheet_name, entries), "excel"
