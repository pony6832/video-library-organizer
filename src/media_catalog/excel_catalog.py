from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
import os
from pathlib import Path
import time
from typing import Callable, Iterable
import uuid
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell import WriteOnlyCell

from .models import MediaRecord, Status


CATALOG_HEADERS = (
    "狀態",
    "檔名",
    "完整路徑",
    "媒體類型",
    "內容描述",
    "重點",
    "關鍵字",
    "拍攝時間",
    "處理時間",
    "Markdown 路徑",
    "備份路徑",
    "錯誤原因",
)

_STATUS_LABELS = {
    Status.PENDING: "待處理",
    Status.PROCESSING: "處理中",
    Status.ANALYZED: "待確認",
    Status.COMPLETED: "完成",
    Status.SKIPPED: "略過",
    Status.FAILED: "失敗",
    Status.MISSING: "來源已移除",
}

_MEDIA_LABELS = {
    "image/jpeg": "照片 (JPEG)",
    "image/png": "照片 (PNG)",
    "image/webp": "照片 (WebP)",
    "image/heic": "照片 (HEIC)",
    "video/mp4": "影片 (MP4)",
    "video/quicktime": "影片 (MOV)",
    "video/x-matroska": "影片 (MKV)",
    "video/webm": "影片 (WebM)",
}


# Excel supports roughly 66,530 hyperlinks per worksheet; past that it asks to
# repair the workbook. Stay well below the limit and leave the path as text.
MAX_HYPERLINKS = 60_000

# Columns people may edit by hand (in Excel or the media library panel). The
# workbook keeps what we last wrote in a hidden baseline sheet; a cell that no
# longer matches its baseline was edited by a person and survives rebuilds.
EDITABLE_HEADERS = ("狀態", "內容描述", "重點", "關鍵字", "拍攝時間")
BASELINE_SHEET = "_organizer_baseline"
# Deliberately not 完整路徑/檔名: catalog readers find the header row by those.
_BASELINE_HEADERS = ("path_key", "status", "description", "highlights", "keywords", "shot_time")
_FINISHED_LABELS = frozenset(("完成", "待確認", "已審核"))
_OWN_STATUS_LABELS = frozenset(_STATUS_LABELS.values())
# Rebuild again when another program saved the workbook mid-rebuild.
_REBUILD_ATTEMPTS = 3

# Leftover temporary workbooks from a killed writer are swept after this age.
_STALE_TEMPORARY_SECONDS = 3600


class ReviewedPathsError(RuntimeError):
    pass


def read_reviewed_paths(excel_path: Path) -> set[str]:
    source = Path(excel_path)
    if not source.is_file():
        return set()
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException):
        return set()
    try:
        sheet = workbook["媒體清冊"]
        headers = {
            cell.value: index
            for index, cell in enumerate(sheet[1], start=1)
            if isinstance(cell.value, str)
        }
        status_column = headers.get("狀態")
        path_column = headers.get("完整路徑")
        if status_column is None or path_column is None:
            return set()
        return {
            str(row[path_column - 1])
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if len(row) >= max(status_column, path_column)
            and row[status_column - 1] == "已審核"
            and row[path_column - 1]
        }
    finally:
        workbook.close()


def read_reviewed_paths_strict(excel_path: Path) -> set[str]:
    source = Path(excel_path)
    if not source.is_file():
        raise ReviewedPathsError("找不到 Excel 媒體清冊")
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except PermissionError:
        raise
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as error:
        raise ReviewedPathsError("Excel 媒體清冊無法讀取") from error
    try:
        if "媒體清冊" not in workbook.sheetnames:
            raise ReviewedPathsError("Excel 缺少媒體清冊工作表")
        sheet = workbook["媒體清冊"]
        rows = sheet.iter_rows(values_only=True)
        headers = {
            value: index
            for index, value in enumerate(next(rows, ()))
            if isinstance(value, str)
        }
        status_column = headers.get("狀態")
        path_column = headers.get("完整路徑")
        if status_column is None or path_column is None:
            raise ReviewedPathsError("Excel 缺少狀態或完整路徑欄位")
        return {
            str(row[path_column])
            for row in rows
            if len(row) > max(status_column, path_column)
            and row[status_column] == "已審核"
            and row[path_column]
        }
    except (BadZipFile, ParseError, KeyError, ValueError) as error:
        raise ReviewedPathsError("Excel 媒體清冊內容損壞，請先修復或還原備份") from error
    finally:
        workbook.close()


def _path_key(value: object) -> str:
    return str(Path(str(value)).resolve()).casefold()


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.isoformat()
    text = str(value).replace("\r\n", "\n")
    return text if text.strip() else None


def _timestamp(value: str | None) -> float | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.timestamp()


def _file_stamp(path: Path) -> tuple[int, int] | None:
    try:
        stat = path.stat()
    except FileNotFoundError:
        return None
    return stat.st_mtime_ns, stat.st_size


@dataclass
class _ExistingCatalog:
    """What a rebuild must carry over from the workbook it replaces."""

    stamp: tuple[int, int] | None = None
    reviewed: set[str] = field(default_factory=set)
    # path key -> {header: text} for cells edited since our last write.
    manual: dict[str, dict[str, str | None]] = field(default_factory=dict)
    # path key -> {header: text} of finished rows in a workbook without baseline.
    legacy: dict[str, dict[str, str | None]] = field(default_factory=dict)
    # path key -> {header: text} of filled cells in rows we did not write.
    foreign: dict[str, dict[str, str | None]] = field(default_factory=dict)
    extra_headers: list[str] = field(default_factory=list)
    extra_values: dict[str, dict[str, object]] = field(default_factory=dict)

    @property
    def saved_at(self) -> float | None:
        return self.stamp[0] / 1e9 if self.stamp else None


def _read_existing(source: Path) -> _ExistingCatalog:
    """Read review marks, hand edits and added columns before a rebuild.

    Fails like ``read_reviewed_paths_strict`` so a damaged workbook is never
    replaced and a locked one is retried later.
    """
    existing = _ExistingCatalog(stamp=_file_stamp(source))
    if existing.stamp is None:
        return existing
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except PermissionError:
        raise
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as error:
        raise ReviewedPathsError("Excel 媒體清冊無法讀取") from error
    try:
        if "媒體清冊" not in workbook.sheetnames:
            raise ReviewedPathsError("Excel 缺少媒體清冊工作表")
        baseline: dict[str, tuple[str | None, ...]] | None = None
        if BASELINE_SHEET in workbook.sheetnames:
            baseline = {}
            for row in workbook[BASELINE_SHEET].iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    cells = tuple(_cell_text(value) for value in row[1:])
                    baseline[str(row[0])] = cells + (None,) * (len(EDITABLE_HEADERS) - len(cells))
        rows = workbook["媒體清冊"].iter_rows(values_only=True)
        headers = {
            value: index
            for index, value in enumerate(next(rows, ()))
            if isinstance(value, str) and value.strip()
        }
        status_column = headers.get("狀態")
        path_column = headers.get("完整路徑")
        if status_column is None or path_column is None:
            raise ReviewedPathsError("Excel 缺少狀態或完整路徑欄位")
        editable = [(position, header) for position, header in enumerate(EDITABLE_HEADERS) if header in headers]
        existing.extra_headers = [header for header in headers if header not in CATALOG_HEADERS]
        for row in rows:
            if len(row) <= path_column or not row[path_column]:
                continue
            key = _path_key(row[path_column])
            status = row[status_column] if len(row) > status_column else None
            if status == "已審核":
                existing.reviewed.add(key)
            cells = {
                header: _cell_text(row[headers[header]]) if len(row) > headers[header] else None
                for _, header in editable
            }
            if baseline is not None:
                # Blank cells and our own status labels are left over from an
                # interrupted run, not edits; the catalog fills them in again.
                # A row missing from the baseline was not written by us (the
                # media library panel appends rows for videos it shows before
                # our next rebuild), so its filled cells are compared with the
                # catalog instead.
                written = baseline.get(key)
                edited = {
                    header: cells[header]
                    for position, header in editable
                    if cells[header] is not None
                    and (written is None or cells[header] != written[position])
                    and not (header == "狀態" and cells[header] in _OWN_STATUS_LABELS)
                }
                if edited:
                    (existing.manual if written is not None else existing.foreign)[key] = edited
            elif status in _FINISHED_LABELS:
                # Status labels change as analysis runs; only content counts.
                existing.legacy[key] = {
                    header: value for header, value in cells.items() if header != "狀態"
                }
            extras = {
                header: row[headers[header]]
                for header in existing.extra_headers
                if len(row) > headers[header] and row[headers[header]] is not None
            }
            if extras:
                existing.extra_values[key] = extras
        return existing
    except (BadZipFile, ParseError, KeyError, ValueError) as error:
        raise ReviewedPathsError("Excel 媒體清冊內容損壞，請先修復或還原備份") from error
    finally:
        workbook.close()


def _generated_cells(record: MediaRecord) -> dict[str, str | None]:
    return {
        "狀態": _STATUS_LABELS[record.status],
        "內容描述": _cell_text(record.description),
        "重點": _cell_text("；".join(record.highlights)),
        "關鍵字": _cell_text("、".join(record.keywords)),
        "拍攝時間": None,
    }


def _hand_edits(
    existing: _ExistingCatalog,
    key: str,
    record: MediaRecord,
    generated: dict[str, str | None],
) -> dict[str, str | None]:
    if key in existing.manual:
        return existing.manual[key]
    if key in existing.foreign:
        return {
            header: value
            for header, value in existing.foreign[key].items()
            if value != generated[header]
        }
    legacy = existing.legacy.get(key)
    if not legacy:
        return {}
    # A workbook from before the baseline sheet: a finished row that differs
    # from the catalog was edited by hand, unless the record changed after the
    # workbook was last saved (then the newer analysis wins). Blank cells are
    # not trusted as edits here because older rebuilds could leave them blank.
    changed_at = _timestamp(record.updated_at)
    saved_at = existing.saved_at
    if changed_at is None or saved_at is None or changed_at > saved_at:
        return {}
    return {
        header: value
        for header, value in legacy.items()
        if value is not None and value != generated[header]
    }


def write_excel(records: Iterable[MediaRecord], output_path: Path) -> Path:
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    # A stream can be read only once, so only a list gets a second rebuild.
    attempts = _REBUILD_ATTEMPTS if isinstance(records, (list, tuple)) else 1
    for attempt in range(1, attempts + 1):
        existing = _read_existing(destination)
        if _rebuild(records, destination, existing, last_attempt=attempt == attempts):
            break
    return destination


def _rebuild(
    records: Iterable[MediaRecord],
    destination: Path,
    existing: _ExistingCatalog,
    *,
    last_attempt: bool,
) -> bool:
    """Write the workbook; False when someone saved it meanwhile (retry)."""
    reviewed_paths = existing.reviewed
    extra_headers = existing.extra_headers

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet()
    sheet.title = "媒體清冊"
    widths = (12, 28, 64, 18, 44, 36, 30, 20, 24, 54, 54, 40) + (18,) * len(extra_headers)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    baseline_sheet = workbook.create_sheet(BASELINE_SHEET)
    baseline_sheet.sheet_state = "hidden"
    baseline_sheet.append(_BASELINE_HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header = []
    for value in CATALOG_HEADERS + tuple(extra_headers):
        cell = WriteOnlyCell(sheet, value=value)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        header.append(cell)
    sheet.append(header)

    row_count = 0
    hyperlink_count = 0
    for record in records:
        # Catalog paths are stored resolved; avoid a resolve() syscall per row.
        source_path = (
            record.path if record.path.is_absolute() else record.path.resolve()
        )
        path_text = str(source_path)
        key = path_text.casefold()
        generated = _generated_cells(record)
        edits = _hand_edits(existing, key, record, generated)
        shown = {**generated, **edits}
        if "狀態" not in edits and key in reviewed_paths:
            shown["狀態"] = "已審核"
        extras = existing.extra_values.get(key, {})
        values = (
                shown["狀態"],
                record.path.name,
                path_text,
                _MEDIA_LABELS.get(record.media_type, record.media_type),
                shown["內容描述"],
                shown["重點"],
                shown["關鍵字"],
                shown["拍攝時間"],
                (
                    record.updated_at
                    if record.status in {Status.ANALYZED, Status.COMPLETED}
                    else None
                ),
                str(record.markdown_path) if record.markdown_path else None,
                str(record.backup_path) if record.backup_path else None,
                record.error,
            ) + tuple(extras.get(header) for header in extra_headers)
        # Filenames and model output are data, even when they start with '='.
        row = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            row.append(cell)
        path_cell = row[2]
        if hyperlink_count < MAX_HYPERLINKS and source_path.is_file():
            path_cell.hyperlink = source_path.as_uri()
            path_cell.style = "Hyperlink"
            hyperlink_count += 1
        sheet.append(row)
        baseline_row = []
        for value in (key,) + tuple(generated[header] for header in EDITABLE_HEADERS):
            cell = WriteOnlyCell(baseline_sheet, value=value)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            baseline_row.append(cell)
        baseline_sheet.append(baseline_row)
        row_count += 1

    if row_count:
        review_validation = DataValidation(
            type="list",
            formula1='"待確認,已審核"',
            allow_blank=False,
            showDropDown=False,
        )
        review_validation.errorTitle = "無效狀態"
        review_validation.error = "請從下拉選單選擇待確認或已審核。"
        review_validation.showErrorMessage = True
        sheet.data_validations.append(review_validation)
        review_validation.add(f"A2:A{row_count + 1}")

    last_column = get_column_letter(len(CATALOG_HEADERS) + len(extra_headers))
    sheet.auto_filter.ref = f"A1:{last_column}{row_count + 1}"
    _sweep_stale_temporaries(destination)
    temporary = destination.with_name(
        f".{destination.stem}.{uuid.uuid4().hex}.tmp{destination.suffix}"
    )
    try:
        workbook.save(temporary)
        workbook.close()
        # The media library panel may have written an edit while we rebuilt;
        # replacing now would drop it, so read the workbook again first.
        if not last_attempt and _file_stamp(destination) != existing.stamp:
            return False
        os.replace(temporary, destination)
    finally:
        workbook.close()
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass
    return True


def _sweep_stale_temporaries(destination: Path) -> None:
    cutoff = time.time() - _STALE_TEMPORARY_SECONDS
    pattern = f".{destination.stem}.*.tmp{destination.suffix}"
    for leftover in destination.parent.glob(pattern):
        try:
            if leftover.stat().st_mtime < cutoff:
                leftover.unlink()
        except OSError:
            pass


def try_write_excel(
    records: Iterable[MediaRecord],
    output_path: Path,
    *,
    writer: Callable[[Iterable[MediaRecord], Path], Path] = write_excel,
) -> bool:
    """Write the catalog, returning False when Excel holds the file open.

    SQLite stays the source of truth, so a locked workbook only delays the
    refresh; the next successful write rebuilds it completely.
    """
    try:
        writer(records, output_path)
    except PermissionError:
        if not Path(output_path).exists():
            raise
        return False
    return True


class ExcelCheckpoint:
    """Refresh Excel every ``every_items`` items, but not more often than
    ``min_interval_seconds``; rewriting a large catalog every 100 items makes
    the total cost quadratic."""

    def __init__(
        self,
        write: Callable[[], object],
        *,
        every_items: int = 100,
        min_interval_seconds: float = 60.0,
        clock: Callable[[], float] = time.monotonic,
    ) -> None:
        self._write = write
        self._every_items = max(1, every_items)
        self._min_interval_seconds = max(0.0, min_interval_seconds)
        self._clock = clock
        self._last_write: float | None = None

    def __call__(self, count: int) -> None:
        if count % self._every_items:
            return
        now = self._clock()
        if (
            self._last_write is not None
            and now - self._last_write < self._min_interval_seconds
        ):
            return
        self._write()
        self._last_write = self._clock()
