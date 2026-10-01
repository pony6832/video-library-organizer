"""Real, local-only release acceptance. Creates a new QA directory, never cleans."""
import argparse
import hashlib
import json
import os
from pathlib import Path
import stat
import subprocess
import sys

from PIL import Image

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT / 'src'))
from media_catalog.database import CatalogDatabase
from media_catalog.workspace import MediaWorkspace
from media_catalog.setup_environment import app_data_root


def run(command, *, env, timeout=120):
    result = subprocess.run([str(item) for item in command], env=env, capture_output=True,
                            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0), timeout=timeout)
    if result.returncode:
        raise RuntimeError(f'QA command failed with exit {result.returncode}: {Path(command[0]).name}')
    return result


def check_bundled_viewer(installed, fixture, records, env):
    """The installer's media library panel opens the analysed folder and reads its database."""
    import time
    import urllib.request
    pin = json.loads((PROJECT / 'packaging' / 'viewer.json').read_text(encoding='utf-8'))
    viewer = installed / 'viewer' / 'VideoLibraryViewer.exe'
    assert hashlib.sha256(viewer.read_bytes()).hexdigest() == pin['assets']['VideoLibraryViewer.exe']
    assert (installed / 'viewer' / 'LAN-share.bat').is_file()
    port = 18790
    process = subprocess.Popen([str(viewer), '--port', str(port), '--no-browser', '--import', str(fixture)],
                               cwd=viewer.parent, env=env, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                               stderr=subprocess.DEVNULL, creationflags=subprocess.CREATE_NO_WINDOW)

    def api(path):
        with urllib.request.urlopen(f'http://127.0.0.1:{port}{path}', timeout=10) as response:
            return json.loads(response.read())
    try:
        for _ in range(60):
            try:
                batches = api('/api/sync')['batches']
                if batches:
                    break
            except OSError:
                pass
            time.sleep(0.5)
        else:
            raise RuntimeError('bundled viewer did not import the folder')
        assert api('/api/status')['version'] == pin['version']
        assert batches[0]['organizer_db'], 'viewer did not link the analysis database'
        items = api('/api/media')['items']
        assert sorted(i['filename'] for i in items) == sorted(r.path.name for r in records)
        assert all(i['db_values'] for i in items)
    finally:
        subprocess.run(['taskkill', '/F', '/T', '/PID', str(process.pid)], capture_output=True,
                       creationflags=subprocess.CREATE_NO_WINDOW)
    return pin['version']


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('release', type=Path)
    parser.add_argument('--qa-id', required=True)
    parser.add_argument('--analyze', action='store_true')
    parser.add_argument('--uninstall', action='store_true')
    args = parser.parse_args()
    if not args.qa_id.replace('-', '').isalnum():
        raise ValueError('qa-id must be a simple new identifier')
    qa = PROJECT / 'build' / 'desktop-acceptance' / args.qa_id
    for parent in [qa, *qa.parents]:
        if parent.exists() and parent.lstat().st_file_attributes & stat.FILE_ATTRIBUTE_REPARSE_POINT:
            raise ValueError('QA path cannot cross reparse points')
    qa.mkdir(parents=True, exist_ok=False)
    release = args.release.resolve(strict=True)
    if not release.is_relative_to(PROJECT / 'dist'):
        raise ValueError('release must be inside project dist')
    env = {k: v for k, v in os.environ.items() if not any(token in k.upper() for token in ('GEMINI', 'GOOGLE_API', 'OPENAI_API', 'ANTHROPIC_API'))}
    env['PYTHONIOENCODING'] = 'utf-8'
    portable = release / 'MediaCatalogVideoDesktop' / 'MediaCatalogVideoDesktop.exe'
    run([portable, '--help'], env=env)
    # A missing optional runtime is a valid diagnostics result, not a crash.
    diagnostic = qa / 'portable-diagnostics.json'
    result = subprocess.run([str(portable), '--diagnostics', str(diagnostic)], env=env,
                            creationflags=subprocess.CREATE_NO_WINDOW, timeout=60)
    assert result.returncode in (0, 1) and 'checks' in json.loads(diagnostic.read_text(encoding='utf-8'))
    installed = qa / 'installed'
    run([release / 'VideoLibraryOrganizer-Setup.exe', '/VERYSILENT', '/SUPPRESSMSGBOXES',
         '/NORESTART', '/QAINSTALL=1', '/NOICONS', '/TASKS=', f'/DIR={installed}',
         f'/LOG={qa / "install.log"}'], env=env)
    exe = installed / 'MediaCatalogVideoDesktop.exe'
    assert hashlib.sha256(exe.read_bytes()).digest() == hashlib.sha256(portable.read_bytes()).digest()
    run([exe, '--help'], env=env)
    fixture = qa / 'synthetic-media'
    fixture.mkdir()
    video = fixture / 'sample-video.mp4'
    photo = fixture / 'skip-photo.png'
    run(['ffmpeg', '-hide_banner', '-loglevel', 'error', '-f', 'lavfi', '-i',
         'testsrc2=size=320x180:rate=12', '-t', '3', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', str(video)], env=env)
    Image.new('RGB', (32, 32), 'blue').save(photo)
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (video, photo)}
    run([exe, '--catalog', 'start', fixture, '--video-only'], env=env)
    workspace = MediaWorkspace.from_root(fixture)
    records = CatalogDatabase(workspace.database_path).list_records()
    assert [r.path.name for r in records] == ['sample-video.mp4']
    if args.analyze:
        run([exe, '--worker', 'analyze-all', fixture, '--skill-root', app_data_root(), '--video-only'], env=env, timeout=900)
        records = CatalogDatabase(workspace.database_path).list_records()
        assert all(r.status.value == 'analyzed' and r.description for r in records)
    from openpyxl import load_workbook
    workbook = load_workbook(workspace.excel_path)
    sheet = workbook.active
    assert sheet.max_row == 2 and sheet.max_column == 12 and sheet['C2'].hyperlink
    workbook.close()
    assert before == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (video, photo)}
    viewer_version = check_bundled_viewer(installed, fixture, records, env)
    receipt = {'installed_exe': str(exe), 'fixture': str(fixture), 'local_analysis': args.analyze,
               'video_records': len(records), 'photos_in_catalog': 0, 'original_sha256': before,
               'excel_rows': 2, 'excel_columns': 12, 'file_hyperlink': True, 'paid_api_calls': False,
               'bundled_viewer': f'{viewer_version}: linked the analysed folder'}
    if args.uninstall:
        runtime = app_data_root() / '.tools/mcp-video-analyzer/node_modules/mcp-video-analyzer/package.json'
        runtime_hash = hashlib.sha256(runtime.read_bytes()).hexdigest() if runtime.exists() else None
        result_hash = hashlib.sha256(workspace.database_path.read_bytes()).hexdigest()
        run([installed / 'unins000.exe', '/VERYSILENT', '/SUPPRESSMSGBOXES', '/NORESTART'], env=env)
        assert not exe.exists()
        assert not (installed / 'viewer' / 'VideoLibraryViewer.exe').exists()
        assert workspace.excel_path.is_file()
        assert result_hash == hashlib.sha256(workspace.database_path.read_bytes()).hexdigest()
        assert before == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (video, photo)}
        if runtime_hash:
            assert runtime_hash == hashlib.sha256(runtime.read_bytes()).hexdigest()
        receipt['qa_uninstall'] = 'passed; originals, results and user runtime retained'
    (qa / 'acceptance.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(receipt, ensure_ascii=True))


if __name__ == '__main__':
    main()
